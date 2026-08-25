import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx.oxml.ns import qn

doc = docx.Document('D:/project/Форматер/Шаблон записки со стилями_Эталон.docx')

# Check specific paragraphs - the ones with direct left alignment on Normal style
# From earlier investigation: para 14, 17, 18, etc.
# But let's find the EXACT ones that audit flags
# Report says page ~2 and page ~139

# Check paragraph 14 (page ~2 area) and paragraph 1295 (page ~139 area)
for idx in [14, 1295]:
    if idx >= len(doc.paragraphs):
        continue
    p = doc.paragraphs[idx]
    print(f'=== Para {idx} ===')
    print(f'  text: {p.text[:50]!r}')
    print(f'  style: {p.style.name}')

    pPr = p._element.pPr
    if pPr is not None:
        jc_in_pPr = pPr.find(qn('w:jc'))
        print(f'  w:jc in pPr: {jc_in_pPr is not None}, val={jc_in_pPr.get(qn("w:val")) if jc_in_pPr is not None else "N/A"}')
        print(f'  pPr XML: {pPr.xml}')
    else:
        print('  pPr: None')

    style_align = p.style.paragraph_format.alignment
    direct_align = p.paragraph_format.alignment
    print(f'  style alignment: {style_align}')
    print(f'  paragraph_format.alignment: {direct_align}')
    print()
