#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
StyleApplier - применение стилей из DOCX-шаблона к целевому документу.
Аналог --reference-doc в Pandoc: копирует стили, настройки страницы,
колонтитулы из шаблона в целевой документ.

Использование:
    from src.apply.style_applier import StyleApplier
    applier = StyleApplier("template.docx")
    applier.apply("target.docx", "output.docx")

Или через CLI:
    python -m src.apply.style_applier --template template.docx --target target.docx --output output.docx
"""

import os
import copy
import logging
import argparse
from typing import Dict, Any, Optional, List, Tuple
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH

logger = logging.getLogger(__name__)

EMU_TO_CM = 1 / 360000


class StyleApplier:
    """
    Применяет стили из DOCX-шаблона к целевому документу.

    Поддерживаемые операции:
    - Копирование параграфных стилей
    - Копирование настроек страницы (поля, размер бумаги)
    - Копирование колонтитулов
    - Очистка прямого форматирования
    """

    STYLES_TO_SKIP = {
        'Default Paragraph Font',
        'Normal',
        'Обычный',
        'No Spacing',
        'Body Text',
    }

    def __init__(self, template_path: str):
        """
        :param template_path: путь к DOCX-шаблону (reference document).
        """
        self.template_path = template_path
        self._template_doc = None

    def _load_template(self):
        """Загружает шаблон (ленивая загрузка)."""
        if self._template_doc is None:
            if not os.path.exists(self.template_path):
                raise FileNotFoundError(f"Файл шаблона не найден: {self.template_path}")
            self._template_doc = Document(self.template_path)

    def _unload_template(self):
        """Освобождает память шаблона."""
        self._template_doc = None

    def apply(
        self,
        target_path: str,
        output_path: Optional[str] = None,
        copy_styles: bool = True,
        copy_page_setup: bool = True,
        copy_headers_footers: bool = True,
        clear_direct_formatting: bool = False,
        styles_filter: Optional[List[str]] = None,
    ) -> Dict[str, Any]:
        """
        Применяет стили из шаблона к целевому документу.

        :param target_path: путь к целевому документу.
        :param output_path: путь для сохранения результата. Если None, перезаписывает target.
        :param copy_styles: копировать стили (True по умолчанию).
        :param copy_page_setup: копировать настройки страницы (True по умолчанию).
        :param copy_headers_footers: копировать колонтитулы (True по умолчанию).
        :param clear_direct_formatting: очистить прямое форматирование перед применением.
        :param styles_filter: список стилей для копирования. Если None, все.
        :return: статистика {'styles_applied': int, 'pages_configured': bool, 'headers_footers_copied': bool}
        """
        self._load_template()

        if not os.path.exists(target_path):
            raise FileNotFoundError(f"Целевой файл не найден: {target_path}")

        target_doc = Document(target_path)

        stats = {
            'styles_applied': 0,
            'pages_configured': False,
            'headers_footers_copied': False,
        }

        if clear_direct_formatting:
            self._clear_all_direct_formatting(target_doc)

        if copy_styles:
            stats['styles_applied'] = self._copy_styles(target_doc, styles_filter)

        if copy_page_setup:
            stats['pages_configured'] = self._copy_page_setup(target_doc)

        if copy_headers_footers:
            stats['headers_footers_copied'] = self._copy_headers_footers(target_doc)

        save_path = output_path if output_path else target_path
        os.makedirs(os.path.dirname(save_path) or '.', exist_ok=True)
        target_doc.save(save_path)

        self._unload_template()

        logger.info(
            f"Стили применены: {stats['styles_applied']} стилей, "
            f"страница: {'да' if stats['pages_configured'] else 'нет'}, "
            f"колонтитулы: {'да' if stats['headers_footers_copied'] else 'нет'}"
        )

        return stats

    # ------------------------------------------------------------------
    # Копирование стилей
    # ------------------------------------------------------------------

    def _copy_styles(self, target_doc, styles_filter: Optional[List[str]] = None) -> int:
        """Копирует стили из шаблона в целевой документ. Возвращает количество."""
        count = 0

        for template_style in self._template_doc.styles:
            if template_style.type is None:
                continue
            try:
                style_type_name = template_style.type.name
            except AttributeError:
                continue

            if style_type_name != 'PARAGRAPH':
                continue

            style_name = template_style.name
            if style_name in self.STYLES_TO_SKIP:
                continue

            if styles_filter and style_name not in styles_filter:
                continue

            try:
                self._copy_single_style(target_doc, template_style)
                count += 1
                logger.debug(f"Скопирован стиль: {style_name}")
            except Exception as e:
                logger.warning(f"Не удалось скопировать стиль '{style_name}': {e}")

        return count

    def _copy_single_style(self, target_doc, template_style):
        """Копирует один стиль из шаблона в целевой документ."""
        style_name = template_style.name

        target_style = self._ensure_style_exists(target_doc, style_name, template_style)

        self._copy_font(template_style, target_style)
        self._copy_paragraph_format(template_style, target_style)

    def _ensure_style_exists(self, target_doc, style_name: str, template_style):
        """Обеспечивает наличие стиля в целевом документе."""
        try:
            return target_doc.styles[style_name]
        except KeyError:
            pass

        try:
            new_style = target_doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        except Exception:
            new_style = target_doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)

        base_name = 'Normal'
        if 'Обычный' in target_doc.styles:
            base_name = 'Обычный'
        try:
            new_style.base_style = target_doc.styles[base_name]
        except Exception:
            pass

        return new_style

    def _copy_font(self, source_style, target_style):
        """Копирует настройки шрифта из source_style в target_style."""
        src_rPr = source_style.element.find(qn('w:rPr'))
        if src_rPr is None:
            return

        tgt_rPr = target_style.element.find(qn('w:rPr'))
        if tgt_rPr is None:
            tgt_rPr = OxmlElement('w:rPr')
            target_style.element.insert(0, tgt_rPr)

        for child in list(tgt_rPr):
            tgt_rPr.remove(child)

        for child in src_rPr:
            tgt_rPr.append(copy.deepcopy(child))

        font = source_style.font
        if font.name:
            target_style.font.name = font.name
        if font.size:
            target_style.font.size = font.size
        if font.bold is not None:
            target_style.font.bold = font.bold
        if font.italic is not None:
            target_style.font.italic = font.italic

    def _copy_paragraph_format(self, source_style, target_style):
        """Копирует настройки абзаца из source_style в target_style."""
        src_pPr = source_style.element.find(qn('w:pPr'))
        if src_pPr is None:
            return

        tgt_pPr = target_style.element.find(qn('w:pPr'))
        if tgt_pPr is None:
            tgt_pPr = OxmlElement('w:pPr')
            target_style.element.append(tgt_pPr)

        for child in list(tgt_pPr):
            tgt_pPr.remove(child)

        for child in src_pPr:
            tgt_pPr.append(copy.deepcopy(child))

        src_fmt = source_style.paragraph_format
        tgt_fmt = target_style.paragraph_format

        if src_fmt.alignment is not None:
            tgt_fmt.alignment = src_fmt.alignment
        if src_fmt.first_line_indent is not None:
            tgt_fmt.first_line_indent = src_fmt.first_line_indent
        if src_fmt.left_indent is not None:
            tgt_fmt.left_indent = src_fmt.left_indent
        if src_fmt.right_indent is not None:
            tgt_fmt.right_indent = src_fmt.right_indent
        if src_fmt.space_before is not None:
            tgt_fmt.space_before = src_fmt.space_before
        if src_fmt.space_after is not None:
            tgt_fmt.space_after = src_fmt.space_after
        if src_fmt.line_spacing is not None:
            tgt_fmt.line_spacing = src_fmt.line_spacing
        if src_fmt.line_spacing_rule is not None:
            tgt_fmt.line_spacing_rule = src_fmt.line_spacing_rule
        if src_fmt.page_break_before is not None:
            tgt_fmt.page_break_before = src_fmt.page_break_before

    # ------------------------------------------------------------------
    # Копирование настроек страницы
    # ------------------------------------------------------------------

    def _copy_page_setup(self, target_doc) -> bool:
        """Копирует настройки страницы из шаблона в целевой документ."""
        if not self._template_doc.sections:
            return False

        template_section = self._template_doc.sections[0]

        for i, target_section in enumerate(target_doc.sections):
            if i > 0 and len(self._template_doc.sections) <= 1:
                break

            src_section = self._template_doc.sections[min(i, len(self._template_doc.sections) - 1)]

            target_section.page_width = src_section.page_width
            target_section.page_height = src_section.page_height
            target_section.orientation = src_section.orientation

            target_section.left_margin = src_section.left_margin
            target_section.right_margin = src_section.right_margin
            target_section.top_margin = src_section.top_margin
            target_section.bottom_margin = src_section.bottom_margin

            target_section.header_distance = src_section.header_distance
            target_section.footer_distance = src_section.footer_distance

            target_section.gutter = src_section.gutter

        logger.debug("Настройки страницы скопированы из шаблона")
        return True

    # ------------------------------------------------------------------
    # Копирование колонтитулов
    # ------------------------------------------------------------------

    def _copy_headers_footers(self, target_doc) -> bool:
        """Копирует колонтитулы из шаблона в целевой документ."""
        if not self._template_doc.sections:
            return False

        template_section = self._template_doc.sections[0]

        for i, target_section in enumerate(target_doc.sections):
            if i > 0 and len(self._template_doc.sections) <= 1:
                break

            src_section = self._template_doc.sections[min(i, len(self._template_doc.sections) - 1)]

            self._copy_header(src_section, target_section)
            self._copy_footer(src_section, target_section)

        logger.debug("Колонтитулы скопированы из шаблона")
        return True

    def _copy_header(self, src_section, tgt_section):
        """Копирует верхний колонтитул."""
        try:
            src_header = src_section.header
        except Exception:
            return

        tgt_header = tgt_section.header
        for para in tgt_header.paragraphs:
            para.clear()

        for src_para in src_header.paragraphs:
            tgt_para = tgt_header.add_paragraph()
            self._copy_paragraph_content(src_para, tgt_para)

            if src_para.alignment is not None:
                tgt_para.alignment = src_para.alignment

            pPr_src = src_para._element.find(qn('w:pPr'))
            if pPr_src is not None:
                pPr_tgt = tgt_para._element.find(qn('w:pPr'))
                if pPr_tgt is None:
                    pPr_tgt = OxmlElement('w:pPr')
                    tgt_para._element.insert(0, pPr_tgt)
                jc_src = pPr_src.find(qn('w:jc'))
                if jc_src is not None:
                    jc_tgt = pPr_tgt.find(qn('w:jc'))
                    if jc_tgt is None:
                        jc_tgt = OxmlElement('w:jc')
                        pPr_tgt.append(jc_tgt)
                    jc_tgt.set(qn('w:val'), jc_src.get(qn('w:val')))

        tgt_header.is_linked_to_previous = False

    def _copy_footer(self, src_section, tgt_section):
        """Копирует нижний колонтитул."""
        try:
            src_footer = src_section.footer
        except Exception:
            return

        tgt_footer = tgt_section.footer
        for para in tgt_footer.paragraphs:
            para.clear()

        for src_para in src_footer.paragraphs:
            tgt_para = tgt_footer.add_paragraph()
            self._copy_paragraph_content(src_para, tgt_para)

            if src_para.alignment is not None:
                tgt_para.alignment = src_para.alignment

        tgt_footer.is_linked_to_previous = False

    def _copy_paragraph_content(self, src_para, tgt_para):
        """Копирует содержимое параграфа (runs, поля, разметку)."""
        for run in src_para.runs:
            tgt_run = tgt_para.add_run()
            tgt_run.text = run.text

            if run.font.name:
                tgt_run.font.name = run.font.name
            if run.font.size:
                tgt_run.font.size = run.font.size
            if run.font.bold is not None:
                tgt_run.font.bold = run.font.bold
            if run.font.italic is not None:
                tgt_run.font.italic = run.font.italic

            try:
                color = run.font.color.rgb
                if color:
                    tgt_run.font.color.rgb = color
            except Exception:
                pass

            rPr_src = run._element.find(qn('w:rPr'))
            if rPr_src is not None:
                rFonts_src = rPr_src.find(qn('w:rFonts'))
                if rFonts_src is not None:
                    rFonts_tgt = tgt_run._element.get_or_add_rPr().get_or_add_rFonts()
                    for attr in ['w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs']:
                        val = rFonts_src.get(qn(attr))
                        if val:
                            rFonts_tgt.set(qn(attr), val)

            self._copy_runs_xml(run, tgt_run)

        for elem in src_para._element:
            tag = elem.tag.split('}')[-1] if '}' in elem.tag else elem.tag
            if tag not in ('pPr', 'r'):
                tgt_para._element.append(copy.deepcopy(elem))

    def _copy_runs_xml(self, src_run, tgt_run):
        """Копирует XML-структуру run (включая поля и сложные элементы)."""
        for elem in src_run._element:
            tag = elem.tag.split('}')[-1] if '}' in elem.tag else elem.tag
            if tag == 'rPr':
                continue
            tgt_run._r.append(copy.deepcopy(elem))

    # ------------------------------------------------------------------
    # Очистка прямого форматирования
    # ------------------------------------------------------------------

    def _clear_all_direct_formatting(self, target_doc):
        """Очищает прямое форматирование во всех параграфах."""
        for para in target_doc.paragraphs:
            for run in para.runs:
                run.font.bold = None
                run.font.italic = None
                run.font.underline = None
                run.font.strike = None
                run.font.color.rgb = None
                run.font.highlight_color = None

        logger.debug("Прямое форматирование очищено")

    # ------------------------------------------------------------------
    # Сравнение документов
    # ------------------------------------------------------------------

    def diff(self, target_path: str) -> Dict[str, Any]:
        """
        Сравнивает стили шаблона с целевым документом.
        Возвращает словарь с расхождениями.

        :param target_path: путь к целевому документу.
        :return: {'missing_styles': [...], 'different_styles': [...], 'page_setup_diff': bool}
        """
        self._load_template()
        target_doc = Document(target_path)

        result = {
            'missing_styles': [],
            'different_styles': [],
            'page_setup_diff': False,
        }

        template_styles = {}
        for style in self._template_doc.styles:
            if style.type and style.type.name == 'PARAGRAPH':
                template_styles[style.name] = style

        target_styles = {}
        for style in target_doc.styles:
            if style.type and style.type.name == 'PARAGRAPH':
                target_styles[style.name] = style

        for name, tmpl_style in template_styles.items():
            if name in self.STYLES_TO_SKIP:
                continue
            if name not in target_styles:
                result['missing_styles'].append(name)
            else:
                if self._styles_differ(tmpl_style, target_styles[name]):
                    result['different_styles'].append(name)

        if self._template_doc.sections and target_doc.sections:
            tmpl_section = self._template_doc.sections[0]
            tgt_section = target_doc.sections[0]

            margins_differ = (
                abs((tmpl_section.left_margin or 0) - (tgt_section.left_margin or 0)) > 10000 or
                abs((tmpl_section.right_margin or 0) - (tgt_section.right_margin or 0)) > 10000 or
                abs((tmpl_section.top_margin or 0) - (tgt_section.top_margin or 0)) > 10000 or
                abs((tmpl_section.bottom_margin or 0) - (tgt_section.bottom_margin or 0)) > 10000
            )
            result['page_setup_diff'] = margins_differ

        self._unload_template()
        return result

    def _styles_differ(self, style1, style2) -> bool:
        """Проверяет, различаются ли два стиля по ключевым параметрам."""
        if style1.font.name != style2.font.name:
            return True
        if style1.font.size != style2.font.size:
            return True
        if style1.font.bold != style2.font.bold:
            return True
        if style1.font.italic != style2.font.italic:
            return True

        p1 = style1.paragraph_format
        p2 = style2.paragraph_format
        if p1.alignment != p2.alignment:
            return True
        if p1.space_before != p2.space_before:
            return True
        if p1.space_after != p2.space_after:
            return True
        if p1.first_line_indent != p2.first_line_indent:
            return True

        return False


def main():
    """CLI-точка входа."""
    parser = argparse.ArgumentParser(
        description='Применение стилей из DOCX-шаблона к целевому документу'
    )
    parser.add_argument(
        '--template', '-t',
        required=True,
        help='Путь к DOCX-шаблону (reference document)'
    )
    parser.add_argument(
        '--target',
        required=True,
        help='Путь к целевому документу'
    )
    parser.add_argument(
        '--output', '-o',
        help='Путь для сохранения результата (по умолчанию перезаписывает target)'
    )
    parser.add_argument(
        '--no-styles',
        action='store_true',
        help='Не копировать стили'
    )
    parser.add_argument(
        '--no-page-setup',
        action='store_true',
        help='Не копировать настройки страницы'
    )
    parser.add_argument(
        '--no-headers-footers',
        action='store_true',
        help='Не копировать колонтитулы'
    )
    parser.add_argument(
        '--clear-formatting',
        action='store_true',
        help='Очистить прямое форматирование перед применением'
    )
    parser.add_argument(
        '--diff',
        action='store_true',
        help='Показать расхождения между шаблоном и целевым документом'
    )
    parser.add_argument(
        '--verbose', '-v',
        action='store_true',
        help='Подробный вывод'
    )

    args = parser.parse_args()

    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format='%(asctime)s [%(levelname)s] %(message)s'
    )

    applier = StyleApplier(args.template)

    if args.diff:
        result = applier.diff(args.target)
        print(f"\nРасхождения с шаблоном {args.template}:")
        print(f"  Отсутствующие стили: {len(result['missing_styles'])}")
        for name in result['missing_styles']:
            print(f"    - {name}")
        print(f"  Различающиеся стили: {len(result['different_styles'])}")
        for name in result['different_styles']:
            print(f"    - {name}")
        print(f"  Настройки страницы различаются: {'да' if result['page_setup_diff'] else 'нет'}")
        return 0

    stats = applier.apply(
        args.target,
        output_path=args.output,
        copy_styles=not args.no_styles,
        copy_page_setup=not args.no_page_setup,
        copy_headers_footers=not args.no_headers_footers,
        clear_direct_formatting=args.clear_formatting,
    )

    output = args.output or args.target
    print(f"\nСтили применены из: {args.template}")
    print(f"Целевой документ: {args.target}")
    print(f"Результат сохранён: {output}")
    print(f"  Скопировано стилей: {stats['styles_applied']}")
    print(f"  Настройки страницы: {'скопированы' if stats['pages_configured'] else 'без изменений'}")
    print(f"  Колонтитулы: {'скопированы' if stats['headers_footers_copied'] else 'без изменений'}")

    return 0


if __name__ == '__main__':
    exit(main())
