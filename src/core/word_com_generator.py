"""
Word COM генератор документов.
Создаёт DOCX из шаблона через Microsoft Word API.
Fallback на python-docx если Word недоступен.
"""
import os
import shutil
import logging
from typing import List, Tuple, Optional, Dict, Any

logger = logging.getLogger(__name__)

TEMPLATE = os.path.join(
    os.path.dirname(__file__), '..', '..', '..',
    'd:\\project\\Demographics\\Наборы данных\\Миасский ГО\\result\\'
    'ОБРАЗЕЦ_Том II. Материалы по обоснованию_Миасский.docx'
)


def is_word_available() -> bool:
    """Проверяет доступность Word COM."""
    try:
        import win32com.client
        word = win32com.client.Dispatch("Word.Application")
        word.Quit()
        return True
    except Exception:
        return False


class WordComGenerator:
    """
    Генератор документов через Word COM.
    Используется для финальной сборки с сохранением всех стилей шаблона.
    """

    def __init__(self, template_path: Optional[str] = None):
        self.template_path = template_path or TEMPLATE
        self._word = None
        self._doc = None

    def _start_word(self):
        """Запускает Word."""
        import win32com.client
        self._word = win32com.client.Dispatch("Word.Application")
        self._word.Visible = False
        self._word.DisplayAlerts = False

    def _stop_word(self):
        """Останавливает Word."""
        if self._doc:
            try:
                self._doc.Close(SaveChanges=False)
            except Exception:
                pass
            self._doc = None
        if self._word:
            try:
                self._word.Quit()
            except Exception:
                pass
            self._word = None

    def create_document(
        self,
        output_path: str,
        content: List[Tuple[str, Optional[str]]],
        keep_original_styles: bool = True
    ) -> str:
        """
        Создаёт документ из шаблона с заданным содержимым.

        Args:
            output_path: путь для сохранения
            content: список (текст, имя_стиля_или_None)
            keep_original_styles: сохранять стили шаблона

        Returns:
            Путь к созданному файлу
        """
        if not is_word_available():
            return self._fallback_create(output_path, content)

        try:
            self._start_word()
            return self._create_via_com(output_path, content)
        except Exception as e:
            logger.error(f"Word COM ошибка: {e}")
            return self._fallback_create(output_path, content)
        finally:
            self._stop_word()

    def _create_via_com(
        self,
        output_path: str,
        content: List[Tuple[str, Optional[str]]]
    ) -> str:
        """Создаёт документ через Word COM."""
        # Копируем шаблон
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        shutil.copy2(self.template_path, output_path)

        # Открываем
        self._doc = self._word.Documents.Open(output_path)

        # Очищаем содержимое
        self._doc.Content.Delete()

        # Добавляем контент
        for i, (text, style_name) in enumerate(content):
            # Вставляем текст
            rng = self._doc.Content
            rng.Collapse(0)  # wdCollapseEnd
            rng.InsertAfter(text)

            # Применяем стиль
            if style_name:
                try:
                    self._doc.Paragraphs(
                        self._doc.Paragraphs.Count
                    ).Range.Style = style_name
                except Exception as e:
                    logger.warning(f"Стиль '{style_name}' не применён: {e}")

            # Перенос строки (кроме последнего)
            if i < len(content) - 1:
                rng = self._doc.Content
                rng.Collapse(0)
                rng.InsertAfter('\r')

        # Сохраняем
        self._doc.Save()
        self._doc.Close()
        self._doc = None

        return output_path

    def _fallback_create(
        self,
        output_path: str,
        content: List[Tuple[str, Optional[str]]]
    ) -> str:
        """Fallback через python-docx (без Word)."""
        from docx import Document

        shutil.copy2(self.template_path, output_path)
        doc = Document(output_path)

        # Очищаем
        body = doc.element.body
        for child in list(body):
            tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
            if tag in ('p', 'tbl'):
                body.remove(child)

        # Добавляем контент
        for text, style_name in content:
            p = doc.add_paragraph(text)
            if style_name:
                try:
                    p.style = doc.styles[style_name]
                except KeyError:
                    pass

        doc.save(output_path)
        return output_path

    def format_existing_document(
        self,
        doc_path: str,
        style_rules: dict
    ) -> str:
        """
        Переформатирует существующий документ по правилам.

        Args:
            doc_path: путь к DOCX
            style_rules: {текст_паттерн: имя_стиля}

        Returns:
            Путь к файлу
        """
        if not is_word_available():
            return doc_path

        try:
            self._start_word()
            self._doc = self._word.Documents.Open(doc_path)

            for para in self._doc.Paragraphs:
                text = para.Range.Text.strip()
                if not text:
                    continue
                for pattern, style_name in style_rules.items():
                    if pattern in text:
                        try:
                            para.Range.Style = style_name
                        except Exception:
                            pass
                        break

            self._doc.Save()
            self._doc.Close()
            self._doc = None
            return doc_path
        except Exception as e:
            logger.error(f"Ошибка форматирования: {e}")
            return doc_path
        finally:
            self._stop_word()


# ============================================================
# Удобные функции
# ============================================================

def create_document(output_path, content, template=None):
    """Создаёт документ из шаблона."""
    gen = WordComGenerator(template)
    return gen.create_document(output_path, content)


def format_document(doc_path, style_rules):
    """Форматирует существующий документ."""
    gen = WordComGenerator()
    return gen.format_existing_document(doc_path, style_rules)


def apply_template_to_document(src_path, template_path, output_path):
    """
    Применяет шаблон к существующему документу.
    Создаёт НОВЫЙ файл с стилями шаблона и содержимым исходного документа.
    """
    import shutil
    from docx import Document

    # Копируем шаблон
    shutil.copy2(template_path, output_path)

    # Открываем исходный документ
    src_doc = Document(src_path)

    # Открываем целевой документ (шаблон)
    dst_doc = Document(output_path)

    # Очищаем целевой документ
    body = dst_doc.element.body
    for child in list(body):
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
        if tag in ('p', 'tbl'):
            body.remove(child)

    # Копируем абзацы с сохранением стилей
    for para in src_doc.paragraphs:
        if para.text.strip():
            new_para = dst_doc.add_paragraph(para.text)
            # Пытаемся применить стиль из шаблона
            try:
                new_para.style = dst_doc.styles[para.style.name]
            except KeyError:
                pass  # Стиль не найден — оставляем Normal

    # Копируем таблицы
    for table in src_doc.tables:
        rows = len(table.rows)
        cols = len(table.columns)
        new_table = dst_doc.add_table(rows=rows, cols=cols, style='Table Grid')
        for i, row in enumerate(table.rows):
            for j, cell in enumerate(row.cells):
                new_table.rows[i].cells[j].text = cell.text

    dst_doc.save(output_path)
    return output_path
