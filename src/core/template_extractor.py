#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Модуль извлечения стилей из DOCX-шаблона (reference document).
Аналог --reference-doc в Pandoc: из готового документа извлекаются
стили, настройки страницы, колонтитулы и генерируется YAML-конфиг.

Использование:
    python -m src.core.template_extractor --template template.docx --output config.yaml
"""

import os
import re
import yaml
import logging
import argparse
from typing import Dict, Any, Optional, List, Tuple
from docx import Document
from docx.shared import Pt, Cm, Twips, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

logger = logging.getLogger(__name__)

TWIPS_TO_CM = 1 / 567
TWIPS_TO_PT = 1 / 20

ALIGNMENT_MAP = {
    'both': 'justify',
    'left': 'left',
    'center': 'center',
    'right': 'right',
    None: 'left',
}

HEADING_PATTERNS = [
    (r'^РАЗДЕЛ\s+\d+', 'Heading 1'),
    (r'^ГЛАВА\s+\d+', 'Heading 2'),
    (r'^\d+\.\d+\s+', 'Heading 3'),
    (r'^Таблица\s+\d+', 'Caption'),
    (r'^Рисунок\s+\d+', 'Caption'),
    (r'^Примечание', 'Table Note'),
    (r'^\s*[–−-]\s+', 'List Paragraph'),
    (r'^\s*\d+\.[\s)]', 'List Paragraph'),
]


class TemplateExtractor:
    """
    Извлекает стили из DOCX-шаблона и генерирует YAML-конфиг
    для использования в AuditEngine и ApplyOrchestrator.
    """

    def __init__(self, template_path: str):
        self.template_path = template_path
        self.doc = None
        self.styles_config: Dict[str, Any] = {}
        self.page_setup: Dict[str, Any] = {}
        self.headers_footers: Dict[str, Any] = {}
        self.detection_rules: Dict[str, Any] = {}

    def extract(self) -> Dict[str, Any]:
        """
        Основной метод: извлекает всё из шаблона и возвращает конфиг.
        """
        self.doc = Document(self.template_path)

        self._extract_styles()
        self._extract_page_setup()
        self._extract_headers_footers()
        self._extract_detection_rules()

        config = {
            'version': '4.2',
            'name': f'Из шаблона: {os.path.basename(self.template_path)}',
            'description': f'Конфигурация, извлечённая из DOCX-шаблона {os.path.basename(self.template_path)}',
            'styles': self.styles_config,
            'detection_rules': self.detection_rules,
            'page_setup': self.page_setup,
            'headers_footers': self.headers_footers,
        }

        self.doc = None
        return config

    def save_yaml(self, output_path: str, config: Optional[Dict] = None):
        """Сохраняет конфиг в YAML-файл."""
        if config is None:
            config = self.extract()

        os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)

        with open(output_path, 'w', encoding='utf-8') as f:
            yaml.dump(
                config, f,
                allow_unicode=True,
                default_flow_style=False,
                sort_keys=False,
                width=120,
            )
        logger.info(f"Конфигурация сохранена: {output_path}")

    # ------------------------------------------------------------------
    # Извлечение стилей
    # ------------------------------------------------------------------

    def _extract_styles(self):
        """Извлекает все параграфные стили из документа."""
        for style in self.doc.styles:
            if style.type is None:
                continue
            try:
                style_type_name = style.type.name
            except AttributeError:
                continue

            if style_type_name != 'PARAGRAPH':
                continue

            if style.name in ('Default Paragraph Font', 'Normal', 'Обычный'):
                continue

            style_cfg = self._extract_single_style(style)
            if style_cfg:
                self.styles_config[style.name] = style_cfg

        if 'Normal' not in self.styles_config and 'Обычный' not in self.styles_config:
            normal = self.doc.styles['Normal']
            style_cfg = self._extract_single_style(normal)
            if style_cfg:
                self.styles_config['Normal'] = style_cfg

    def _extract_single_style(self, style) -> Optional[Dict[str, Any]]:
        """Извлекает настройки одного стиля."""
        config: Dict[str, Any] = {'enabled': True}

        font_cfg = self._extract_font(style)
        if font_cfg:
            config['font'] = font_cfg

        para_cfg = self._extract_paragraph(style)
        if para_cfg:
            config['paragraph'] = para_cfg

        if not config.get('font') and not config.get('paragraph'):
            return None

        return config

    def _extract_font(self, style) -> Optional[Dict[str, Any]]:
        """Извлекает настройки шрифта из стиля."""
        font = style.font
        cfg: Dict[str, Any] = {}

        name = font.name
        if name:
            cfg['name'] = name

        if font.size:
            cfg['size'] = int(font.size.pt)

        if font.bold is not None:
            cfg['bold'] = bool(font.bold)

        if font.italic is not None:
            cfg['italic'] = bool(font.italic)

        try:
            color = font.color.rgb
            if color:
                cfg['color'] = str(color)
        except Exception:
            pass

        rPr = style.element.find(qn('w:rPr'))
        if rPr is not None:
            rFonts = rPr.find(qn('w:rFonts'))
            if rFonts is not None:
                east_asia = rFonts.get(qn('w:eastAsia'))
                if east_asia:
                    cfg['east_asia'] = east_asia

        if not cfg:
            return None
        return cfg

    def _extract_paragraph(self, style) -> Optional[Dict[str, Any]]:
        """Извлекает настройки абзаца из стиля."""
        pPr = style.element.find(qn('w:pPr'))
        if pPr is None:
            return None

        cfg: Dict[str, Any] = {}

        ind = pPr.find(qn('w:ind'))
        if ind is not None:
            first_line = self._get_int_attr(ind, 'w:firstLine')
            if first_line is not None:
                cfg['first_line_indent_cm'] = round(first_line * TWIPS_TO_CM, 2)

            left = self._get_int_attr(ind, 'w:left')
            if left is not None:
                cfg['left_indent_cm'] = round(left * TWIPS_TO_CM, 2)

        spacing = pPr.find(qn('w:spacing'))
        if spacing is not None:
            before = self._get_int_attr(spacing, 'w:before')
            if before is not None:
                cfg['space_before_pt'] = round(before * TWIPS_TO_PT, 1)

            after = self._get_int_attr(spacing, 'w:after')
            if after is not None:
                cfg['space_after_pt'] = round(after * TWIPS_TO_PT, 1)

            line = self._get_int_attr(spacing, 'w:line')
            line_rule = spacing.get(qn('w:lineRule'))
            if line is not None:
                if line_rule == 'auto':
                    cfg['line_spacing'] = round(line / 240, 2)
                else:
                    cfg['line_spacing'] = round(line * TWIPS_TO_PT, 1)

        jc = pPr.find(qn('w:jc'))
        if jc is not None:
            val = jc.get(qn('w:val'))
            cfg['alignment'] = ALIGNMENT_MAP.get(val, 'left')

        page_break = pPr.find(qn('w:pageBreakBefore'))
        if page_break is not None:
            val = page_break.get(qn('w:val'))
            if val != '0':
                cfg['page_break_before'] = True

        if not cfg:
            return None
        return cfg

    # ------------------------------------------------------------------
    # Извлечение настроек страницы
    # ------------------------------------------------------------------

    def _extract_page_setup(self):
        """Извлекает настройки страницы из первой секции."""
        if not self.doc.sections:
            return

        section = self.doc.sections[0]

        def emu_to_cm(val):
            """Конвертирует EMU (English Metric Units) в сантиметры."""
            if val is None:
                return None
            return round(val / 360000, 2)

        self.page_setup = {
            'portrait': {
                'left_margin_cm': emu_to_cm(section.left_margin) if section.left_margin else 3.0,
                'right_margin_cm': emu_to_cm(section.right_margin) if section.right_margin else 1.5,
                'top_margin_cm': emu_to_cm(section.top_margin) if section.top_margin else 2.0,
                'bottom_margin_cm': emu_to_cm(section.bottom_margin) if section.bottom_margin else 2.0,
            },
            'paper_size': 'A4',
            'orientation_default': 'portrait',
        }

    # ------------------------------------------------------------------
    # Извлечение колонтитулов
    # ------------------------------------------------------------------

    def _extract_headers_footers(self):
        """Извлекает настройки колонтитулов из первой секции."""
        if not self.doc.sections:
            return

        section = self.doc.sections[0]

        header_cfg = self._extract_header_config(section)
        footer_cfg = self._extract_footer_config(section)

        self.headers_footers = {}
        if header_cfg:
            self.headers_footers['header'] = header_cfg
        if footer_cfg:
            self.headers_footers['footer'] = footer_cfg

    def _extract_header_config(self, section) -> Optional[Dict[str, Any]]:
        """Извлекает конфигурацию верхнего колонтитула."""
        try:
            header = section.header
        except Exception:
            return None

        if not header.paragraphs:
            return None

        cfg: Dict[str, Any] = {'enabled': True}

        if section.header_distance:
            cfg['position_from_edge_cm'] = round(section.header_distance / 360000, 2)

        has_page_field = False
        for para in header.paragraphs:
            for run in para.runs:
                for elem in run._r:
                    if elem.tag == qn('w:fldChar'):
                        has_page_field = True
                        break
                if has_page_field:
                    break
            if has_page_field:
                break

        if has_page_field:
            cfg['content'] = 'page_number'

        for para in header.paragraphs:
            if para.alignment is not None:
                align_val = para.alignment
                try:
                    cfg['alignment'] = ALIGNMENT_MAP.get(str(align_val), 'right')
                except Exception:
                    cfg['alignment'] = 'right'
                break

        if header.paragraphs:
            for run in header.paragraphs[0].runs:
                font = run.font
                font_cfg: Dict[str, Any] = {}
                if font.name:
                    font_cfg['name'] = font.name
                if font.size:
                    font_cfg['size'] = int(font.size.pt)
                if font.bold is not None:
                    font_cfg['bold'] = bool(font.bold)
                if font_cfg:
                    cfg['font'] = font_cfg
                break

        return cfg

    def _extract_footer_config(self, section) -> Optional[Dict[str, Any]]:
        """Извлекает конфигурацию нижнего колонтитула."""
        try:
            footer = section.footer
        except Exception:
            return None

        if not footer.paragraphs:
            return None

        cfg: Dict[str, Any] = {'enabled': True}

        if section.footer_distance:
            cfg['position_from_edge_cm'] = round(section.footer_distance / 360000, 2)

        has_section_field = False
        for para in footer.paragraphs:
            for run in para.runs:
                text = run.text
                if 'РАЗДЕЛ' in text or 'СЕКЦИЯ' in text:
                    has_section_field = True
                    break
                for elem in run._r:
                    if elem.tag == qn('w:fldChar'):
                        has_section_field = True
                        break
                if has_section_field:
                    break
            if has_section_field:
                break

        if has_section_field:
            cfg['content'] = 'section_chapter_number'

        for para in footer.paragraphs:
            if para.alignment is not None:
                try:
                    cfg['alignment'] = ALIGNMENT_MAP.get(str(para.alignment), 'center')
                except Exception:
                    cfg['alignment'] = 'center'
                break

        if footer.paragraphs:
            for run in footer.paragraphs[0].runs:
                font = run.font
                font_cfg: Dict[str, Any] = {}
                if font.name:
                    font_cfg['name'] = font.name
                if font.size:
                    font_cfg['size'] = int(font.size.pt)
                if font.bold is not None:
                    font_cfg['bold'] = bool(font.bold)
                if font_cfg:
                    cfg['font'] = font_cfg
                break

        return cfg

    # ------------------------------------------------------------------
    # Извлечение правил определения стилей
    # ------------------------------------------------------------------

    def _extract_detection_rules(self):
        """Анализирует текст параграфов и генерирует regex-правила определения."""
        style_samples: Dict[str, List[str]] = {}

        for para in self.doc.paragraphs:
            text = para.text.strip()
            if not text or len(text) > 100:
                continue

            style_name = para.style.name
            if style_name in ('Normal', 'Обычный', 'Default Paragraph Font'):
                continue

            if style_name not in style_samples:
                style_samples[style_name] = []

            if len(style_samples[style_name]) < 5:
                style_samples[style_name].append(text)

        self.detection_rules = {}
        for style_name, samples in style_samples.items():
            patterns = self._generate_patterns(samples)
            if patterns:
                self.detection_rules[style_name] = patterns

    def _generate_patterns(self, samples: List[str]) -> List[str]:
        """Генерирует regex-паттерны на основе образцов текста."""
        patterns = []

        prefix_groups: Dict[str, int] = {}
        for sample in samples:
            match = re.match(r'^([А-ЯA-Z]+\s*)', sample)
            if match:
                prefix = match.group(1).strip()
                prefix_groups[prefix] = prefix_groups.get(prefix, 0) + 1

        for prefix, count in prefix_groups.items():
            if count >= 2:
                patterns.append(f'^{re.escape(prefix)}\\s+\\d+')
                patterns.append(f'^{re.escape(prefix)} ')

        numeric_pattern = all(re.match(r'^\d+\.\d+', s) for s in samples)
        if numeric_pattern:
            patterns.append(r'^\d+\.\d+\s+')

        if not patterns:
            for sample in samples[:2]:
                escaped = re.escape(sample[:20])
                patterns.append(f'^{escaped}')

        return patterns[:4]

    # ------------------------------------------------------------------
    # Утилиты
    # ------------------------------------------------------------------

    @staticmethod
    def _get_int_attr(element: OxmlElement, attr_name: str) -> Optional[int]:
        """Безопасно извлекает целочисленный атрибут из XML-элемента."""
        val = element.get(qn(attr_name))
        if val is not None:
            try:
                return int(val)
            except ValueError:
                return None
        return None


def main():
    """CLI-точка входа."""
    parser = argparse.ArgumentParser(
        description='Извлечение стилей из DOCX-шаблона и генерация YAML-конфига'
    )
    parser.add_argument(
        '--template', '-t',
        required=True,
        help='Путь к DOCX-шаблону (reference document)'
    )
    parser.add_argument(
        '--output', '-o',
        required=True,
        help='Путь для сохранения YAML-конфига'
    )
    parser.add_argument(
        '--verbose', '-v',
        action='store_true',
        help='Подробный вывод'
    )

    args = parser.parse_args()

    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format='%(asctime)s [%(levelname)s] %(message)s'
    )

    if not os.path.exists(args.template):
        logger.error(f"Файл шаблона не найден: {args.template}")
        return 1

    extractor = TemplateExtractor(args.template)
    config = extractor.extract()
    extractor.save_yaml(args.output, config)

    print(f"\nКонфигурация извлечена из: {args.template}")
    print(f"Сохранена в: {args.output}")
    print(f"Стилей: {len(config.get('styles', {}))}")
    print(f"Правил определения: {len(config.get('detection_rules', {}))}")

    return 0


if __name__ == '__main__':
    exit(main())
