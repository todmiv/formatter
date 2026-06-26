#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
FormatterService — основной сервис форматирования документов.
Отвечает за форматирование по конфигу и шаблону.
"""

import os
import time
import logging
from typing import Optional, Dict, Any

from src.core.config_loader import ConfigLoader
from src.core.content_processor import ContentProcessor
from src.core.models import FormatResult

logger = logging.getLogger(__name__)


class FormatterService:
    """
    Сервис форматирования документов.

    Поддерживает два режима:
    - config-based: YAML-конфиг задаёт правила форматирования
    - template-based: DOCX-шаблон作为 образец стилей
    """

    def __init__(
        self,
        config_path: Optional[str] = None,
        template_path: Optional[str] = None,
    ):
        from src.core.formatter_errors import raise_config_not_found, raise_template_not_found

        self.config_path = config_path
        self.template_path = template_path

        self.config_loader: Optional[ConfigLoader] = None
        self._style_applier = None
        self._content_processor = None

        if config_path:
            if not os.path.exists(config_path):
                raise_config_not_found(config_path)
            self.config_loader = ConfigLoader(config_path)
            self.config_loader.load()
            self._content_processor = ContentProcessor(self.config_loader)
        elif template_path:
            if not os.path.exists(template_path):
                raise_template_not_found(template_path)
            from src.apply.style_applier import StyleApplier
            self._style_applier = StyleApplier(template_path)

    @property
    def mode(self) -> str:
        return 'config' if self.config_loader else 'template'

    def format_document(
        self,
        input_path: str,
        output_path: Optional[str] = None,
        clear_formatting: bool = False,
        process_content: bool = True,
    ) -> FormatResult:
        """
        Форматирует документ.

        :param input_path: путь к входному документу.
        :param output_path: путь для сохранения результата.
        :param clear_formatting: очистить прямое форматирование.
        :param process_content: обработать содержимое.
        :return: результат форматирования.
        """
        from src.core.formatter_errors import raise_file_not_found

        t_start = time.perf_counter()
        output = output_path or self._default_output_path(input_path)

        if not os.path.exists(input_path):
            raise_file_not_found(input_path)

        result = FormatResult(
            input=input_path,
            output=output,
            mode=self.mode,
            success=False,
        )

        try:
            if self._style_applier:
                stats = self._apply_from_template(input_path, output, clear_formatting)
            elif self.config_loader:
                stats = self._apply_from_config(
                    input_path, output, clear_formatting, process_content
                )
            else:
                raise ValueError("Не задан ни конфиг, ни шаблон")

            result.stats = stats
            result.success = True

        except Exception as e:
            logger.error(f"Ошибка форматирования: {e}", exc_info=True)
            result.error = str(e)

        result.duration = time.perf_counter() - t_start
        logger.info(f"Форматирование завершено: {output} ({result.duration:.2f} сек)")

        return result

    def _apply_from_template(
        self,
        input_path: str,
        output_path: str,
        clear_formatting: bool,
    ) -> Dict[str, Any]:
        """Применяет стили из шаблона."""
        stats = self._style_applier.apply(
            input_path,
            output_path,
            copy_styles=True,
            copy_page_setup=True,
            copy_headers_footers=True,
            clear_direct_formatting=clear_formatting,
        )

        if self._content_processor:
            from docx import Document
            doc = Document(output_path)
            transform_stats = self._content_processor.transform_document(doc)
            stats.update(transform_stats)

        return stats

    def _apply_from_config(
        self,
        input_path: str,
        output_path: str,
        clear_formatting: bool,
        process_content: bool,
    ) -> Dict[str, Any]:
        """Применяет стили из конфигурации."""
        from docx import Document
        from src.apply.apply_orchestrator import ApplyOrchestrator

        doc = Document(input_path)
        all_issues = self._scan_for_all_issues(doc)

        if not all_issues:
            logger.info("Проблем не найдено, копируем документ")
            doc.save(output_path)
            return {'styles_applied': 0, 'content_transforms': 0}

        orchestrator = ApplyOrchestrator(self.config_loader)
        stats = orchestrator.apply_fixes(
            input_path, all_issues, output_path,
            apply_styles=True,
            apply_direct_overrides=True,
            clear_direct_formatting=clear_formatting,
        )

        if process_content and self._content_processor:
            doc_out = Document(output_path)
            transform_stats = self._content_processor.transform_document(doc_out)
            stats['content_transforms'] = transform_stats.get('transforms_applied', 0)

        return stats

    def _scan_for_all_issues(self, doc) -> list:
        """Сканирует документ для поиска всех элементов, требующих форматирования."""
        from src.core.audit_engine import AuditEngine

        if not self.config_loader:
            return []

        engine = AuditEngine(self.config_loader)
        issues = engine.scan_document(doc=self._temp_save(doc))

        return [i for i in issues if i.auto_fixable]

    def _temp_save(self, doc) -> str:
        """Временно сохраняет документ для сканирования."""
        import tempfile
        tmp = tempfile.NamedTemporaryFile(suffix='.docx', delete=False)
        tmp.close()
        doc.save(tmp.name)
        return tmp.name

    def _default_output_path(self, input_path: str) -> str:
        """Генерирует путь вывода по умолчанию."""
        base, ext = os.path.splitext(input_path)
        return f"{base}_formatted{ext}"
