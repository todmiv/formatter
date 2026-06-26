#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
DocumentFormatter - высокоуровневый модуль форматирования документов.
Объединяет StyleApplier, ContentProcessor, AuditEngine и ApplyOrchestrator
в единый интерфейс для форматирования DOCX-документов.

Использование:
    from src.core.document_formatter import DocumentFormatter

    # Форматирование по конфигу
    formatter = DocumentFormatter(config_path='configs/active/config.yaml')
    result = formatter.format_document('input.docx', 'output.docx')

    # Форматирование по шаблону
    formatter = DocumentFormatter(template_path='template.docx')
    result = formatter.format_document('input.docx', 'output.docx')

    # Полный цикл: аудит + исправление + валидация
    result = formatter.format_with_audit('input.docx', 'output.docx')

Или через CLI:
    python -m src.core.document_formatter -i input.docx -o output.docx -c config.yaml
    python -m src.core.document_formatter -i input.docx -o output.docx -t template.docx
"""

import os
import time
import logging
import argparse
from typing import Dict, Any, Optional, List
from docx import Document

from src.core.config_loader import ConfigLoader
from src.core.content_processor import ContentProcessor

logger = logging.getLogger(__name__)


class DocumentFormatter:
    """
    Высокоуровневый форматировщик документов.

    Поддерживает два режима работы:
    1. Config-based: YAML-конфиг задаёт правила форматирования
    2. Template-based: DOCX-шаблон作为 образец стилей
    """

    def __init__(
        self,
        config_path: Optional[str] = None,
        template_path: Optional[str] = None,
        verbose: bool = False,
    ):
        """
        :param config_path: путь к YAML-конфигу (для config-based режима).
        :param template_path: путь к DOCX-шаблону (для template-based режима).
        :param verbose: подробный вывод.
        """
        self.config_path = config_path
        self.template_path = template_path
        self.verbose = verbose

        self.config_loader: Optional[ConfigLoader] = None
        self._style_applier = None
        self._content_processor = None

        if config_path:
            self._init_config_mode(config_path)
        elif template_path:
            self._init_template_mode(template_path)

    def _init_config_mode(self, config_path: str):
        """Инициализация config-based режима."""
        from src.core.formatter_errors import raise_config_not_found

        if not os.path.exists(config_path):
            raise_config_not_found(config_path)

        self.config_loader = ConfigLoader(config_path)
        self.config_loader.load()
        self._content_processor = ContentProcessor(self.config_loader)

        logger.info(f"Config-based режим: {config_path}")

    def _init_template_mode(self, template_path: str):
        """Инициализация template-based режима."""
        from src.core.formatter_errors import raise_template_not_found

        if not os.path.exists(template_path):
            raise_template_not_found(template_path)

        from src.apply.style_applier import StyleApplier
        self._style_applier = StyleApplier(template_path)

        logger.info(f"Template-based режим: {template_path}")

    # ------------------------------------------------------------------
    # Основные методы форматирования
    # ------------------------------------------------------------------

    def format_document(
        self,
        input_path: str,
        output_path: Optional[str] = None,
        options: Optional[Dict[str, Any]] = None,
    ) -> Dict[str, Any]:
        """
        Форматирует документ.

        :param input_path: путь к входному документу.
        :param output_path: путь для сохранения результата.
        :param options: дополнительные опции.
        :return: результат форматирования.
        """
        from src.core.formatter_errors import raise_file_not_found

        if not os.path.exists(input_path):
            raise_file_not_found(input_path)

        t_start = time.perf_counter()
        output = output_path or self._default_output_path(input_path)

        options = options or {}
        clean_formatting = options.get('clear_direct_formatting', False)
        process_content = options.get('process_content', True)

        result = {
            'input': input_path,
            'output': output,
            'mode': 'config' if self.config_loader else 'template',
            'stats': {},
            'success': False,
        }

        try:
            if self._style_applier:
                result['stats'] = self._apply_from_template(
                    input_path, output, clean_formatting
                )
            elif self.config_loader:
                result['stats'] = self._apply_from_config(
                    input_path, output, clean_formatting, process_content
                )
            else:
                raise ValueError("Не задан ни конфиг, ни шаблон")

            result['success'] = True

        except Exception as e:
            logger.error(f"Ошибка форматирования: {e}", exc_info=True)
            result['error'] = str(e)

        result['duration'] = time.perf_counter() - t_start
        logger.info(
            f"Форматирование завершено: {output} "
            f"({result['duration']:.2f} сек)"
        )

        return result

    def format_with_audit(
        self,
        input_path: str,
        output_path: Optional[str] = None,
        max_iterations: int = 3,
    ) -> Dict[str, Any]:
        """
        Полный цикл: аудит → исправление → повторный аудит → валидация.

        :param input_path: путь к входному документу.
        :param output_path: путь для сохранения результата.
        :param max_iterations: максимальное количество итераций.
        :return: результат форматирования с отчётом аудита.
        """
        from src.core.formatter_errors import FormatterError, ErrorKind

        if not self.config_loader:
            raise FormatterError(ErrorKind.CONFIG_NOT_FOUND, detail="format_with_audit требует config_path")

        from src.core.audit_engine import AuditEngine
        from src.apply.apply_orchestrator import ApplyOrchestrator

        t_start = time.perf_counter()
        output = output_path or self._default_output_path(input_path)

        result = {
            'input': input_path,
            'output': output,
            'mode': 'config',
            'iterations': [],
            'final_stats': {},
            'content_issues': [],
            'success': False,
        }

        try:
            current_path = input_path
            audit_engine = AuditEngine(self.config_loader)

            for iteration in range(max_iterations):
                t_iter = time.perf_counter()

                issues = audit_engine.scan_document(current_path)

                content_issues = []
                if self._content_processor:
                    doc = Document(current_path)
                    content_issues = self._content_processor.validate_document(doc)
                    result['content_issues'].extend(content_issues)

                if not issues:
                    logger.info(f"Аудит итерации {iteration + 1}: проблем нет")
                    break

                fixable = [i for i in issues if i.auto_fixable]
                if not fixable:
                    logger.info(f"Аудит итерации {iteration + 1}: нет исправляемых проблем")
                    break

                orchestrator = ApplyOrchestrator(self.config_loader)
                stats = orchestrator.apply_fixes(
                    current_path, fixable, output,
                    apply_styles=True,
                    apply_direct_overrides=True,
                )

                iter_result = {
                    'iteration': iteration + 1,
                    'issues_found': len(issues),
                    'issues_fixed': stats['applied'],
                    'issues_failed': stats['failed'],
                    'duration': time.perf_counter() - t_iter,
                }
                result['iterations'].append(iter_result)

                logger.info(
                    f"Итерация {iteration + 1}: найдено {len(issues)}, "
                    f"исправлено {stats['applied']}, ошибок {stats['failed']}"
                )

                current_path = output

                if stats['failed'] == 0:
                    break

            final_audit = audit_engine.scan_document(output)
            result['final_stats'] = {
                'remaining_issues': len(final_audit),
                'total_iterations': len(result['iterations']),
            }

            result['success'] = True

        except Exception as e:
            logger.error(f"Ошибка форматирования с аудитом: {e}", exc_info=True)
            result['error'] = str(e)

        result['duration'] = time.perf_counter() - t_start
        return result

    def batch_format(
        self,
        input_dir: str,
        output_dir: str,
        pattern: str = '*.docx',
        options: Optional[Dict[str, Any]] = None,
    ) -> List[Dict[str, Any]]:
        """
        Пакетное форматирование документов.

        :param input_dir: входная директория.
        :param output_dir: выходная директория.
        :param pattern: паттерн файлов.
        :param options: опции форматирования.
        :return: список результатов.
        """
        import glob

        os.makedirs(output_dir, exist_ok=True)

        search = os.path.join(input_dir, pattern)
        files = glob.glob(search)

        results = []
        for file_path in files:
            file_name = os.path.basename(file_path)
            output_path = os.path.join(output_dir, file_name)

            logger.info(f"Пакетное форматирование: {file_name}")
            result = self.format_document(file_path, output_path, options)
            results.append(result)

        success = sum(1 for r in results if r['success'])
        logger.info(
            f"Пакетное форматирование завершено: "
            f"{success}/{len(results)} успешно"
        )

        return results

    # ------------------------------------------------------------------
    # Методы сравнения и анализа
    # ------------------------------------------------------------------

    def diff(self, input_path: str, template_path: Optional[str] = None) -> Dict[str, Any]:
        """
        Сравнивает документ с шаблоном или конфигом.

        :param input_path: путь к документу.
        :param template_path: путь к шаблону (если не задан в конструкторе).
        :return: результат сравнения.
        """
        from src.core.formatter_errors import FormatterError, ErrorKind

        template = template_path or self.template_path
        if not template:
            raise FormatterError(ErrorKind.TEMPLATE_NOT_FOUND, detail="для сравнения необходим template_path")

        from src.apply.style_applier import StyleApplier
        applier = StyleApplier(template)
        return applier.diff(input_path)

    def extract_config(
        self,
        template_path: str,
        output_path: str,
    ) -> Dict[str, Any]:
        """
        Извлекает конфигурацию из DOCX-шаблона.

        :param template_path: путь к шаблону.
        :param output_path: путь для сохранения YAML.
        :return: извлечённая конфигурация.
        """
        from src.core.template_extractor import TemplateExtractor

        extractor = TemplateExtractor(template_path)
        config = extractor.extract()
        extractor.save_yaml(output_path, config)

        logger.info(f"Конфигурация извлечена: {output_path}")
        return config

    # ------------------------------------------------------------------
    # Приватные методы
    # ------------------------------------------------------------------

    def _apply_from_template(
        self,
        input_path: str,
        output_path: str,
        clear_formatting: bool,
    ) -> Dict[str, Any]:
        """Применяет стили из шаблона."""
        stats = self._style_applier.apply(
            input_path,
            output_path,
            copy_styles=True,
            copy_page_setup=True,
            copy_headers_footers=True,
            clear_direct_formatting=clear_formatting,
        )

        if self._content_processor:
            doc = Document(output_path)
            transform_stats = self._content_processor.transform_document(doc)
            stats.update(transform_stats)

        return stats

    def _apply_from_config(
        self,
        input_path: str,
        output_path: str,
        clear_formatting: bool,
        process_content: bool,
    ) -> Dict[str, Any]:
        """Применяет стили из конфигурации."""
        from src.apply.apply_orchestrator import ApplyOrchestrator

        doc = Document(input_path)
        all_issues = self._scan_for_all_issues(doc)

        if not all_issues:
            logger.info("Проблем не найдено, копируем документ")
            doc.save(output_path)
            return {'styles_applied': 0, 'content_transforms': 0}

        orchestrator = ApplyOrchestrator(self.config_loader)
        stats = orchestrator.apply_fixes(
            input_path, all_issues, output_path,
            apply_styles=True,
            apply_direct_overrides=True,
            clear_direct_formatting=clear_formatting,
        )

        if process_content and self._content_processor:
            doc_out = Document(output_path)
            transform_stats = self._content_processor.transform_document(doc_out)
            stats['content_transforms'] = transform_stats.get('transforms_applied', 0)

        return stats

    def _scan_for_all_issues(self, doc: Document) -> list:
        """Сканирует документ для поиска всех элементов, требующих форматирования."""
        from src.core.audit_engine import AuditEngine

        if not self.config_loader:
            return []

        engine = AuditEngine(self.config_loader)
        issues = engine.scan_document(self._temp_save(doc))

        return [i for i in issues if i.auto_fixable]

    def _temp_save(self, doc: Document) -> str:
        """Временно сохраняет документ для сканирования."""
        import tempfile
        tmp = tempfile.NamedTemporaryFile(suffix='.docx', delete=False)
        tmp.close()
        doc.save(tmp.name)
        return tmp.name

    def _default_output_path(self, input_path: str) -> str:
        """Генерирует путь вывода по умолчанию."""
        base, ext = os.path.splitext(input_path)
        return f"{base}_formatted{ext}"


def main():
    """CLI-точка входа."""
    parser = argparse.ArgumentParser(
        description='Форматирование DOCX-документов по ГОСТ'
    )
    parser.add_argument(
        '-i', '--input',
        required=True,
        help='Входной DOCX-файл'
    )
    parser.add_argument(
        '-o', '--output',
        help='Выходной DOCX-файл (по умолчанию: input_formatted.docx)'
    )
    parser.add_argument(
        '-c', '--config',
        help='Путь к YAML-конфигу'
    )
    parser.add_argument(
        '-t', '--template',
        help='Путь к DOCX-шаблону (reference document)'
    )
    parser.add_argument(
        '--audit',
        action='store_true',
        help='Выполнить полный цикл: аудит + исправление'
    )
    parser.add_argument(
        '--max-iterations',
        type=int,
        default=3,
        help='Максимальное количество итераций (для --audit)'
    )
    parser.add_argument(
        '--clear-formatting',
        action='store_true',
        help='Очистить прямое форматирование'
    )
    parser.add_argument(
        '--diff',
        action='store_true',
        help='Показать расхождения с шаблоном'
    )
    parser.add_argument(
        '--extract-config',
        help='Извлечь конфиг из шаблона и сохранить в YAML'
    )
    parser.add_argument(
        '-v', '--verbose',
        action='store_true',
        help='Подробный вывод'
    )

    args = parser.parse_args()

    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format='%(asctime)s [%(levelname)s] %(message)s'
    )

    config_path = args.config
    template_path = args.template

    if not config_path and not template_path:
        default_config = 'configs/active/config.yaml'
        if os.path.exists(default_config):
            config_path = default_config
            logger.info(f"Используется конфиг по умолчанию: {config_path}")

    formatter = DocumentFormatter(
        config_path=config_path,
        template_path=template_path,
        verbose=args.verbose,
    )

    if args.extract_config:
        if not template_path:
            logger.error("Для --extract-config необходим --template")
            return 1
        formatter.extract_config(template_path, args.extract_config)
        print(f"Конфигурация извлечена: {args.extract_config}")
        return 0

    if args.diff:
        result = formatter.diff(args.input, template_path)
        print("\nРасхождения с шаблоном:")
        print(f"  Отсутствующие стили: {len(result['missing_styles'])}")
        for name in result['missing_styles']:
            print(f"    - {name}")
        print(f"  Различающиеся стили: {len(result['different_styles'])}")
        for name in result['different_styles']:
            print(f"    - {name}")
        page_status = 'различаются' if result['page_setup_diff'] else 'совпадают'
        print(f"  Настройки страницы: {page_status}")
        return 0

    if args.audit:
        result = formatter.format_with_audit(
            args.input,
            args.output,
            max_iterations=args.max_iterations,
        )
    else:
        result = formatter.format_document(
            args.input,
            args.output,
            options={'clear_direct_formatting': args.clear_formatting},
        )

    if result['success']:
        print("\nФорматирование завершено успешно")
        print(f"  Вход: {result['input']}")
        print(f"  Выход: {result['output']}")
        print(f"  Режим: {result['mode']}")
        print(f"  Время: {result['duration']:.2f} сек")

        if 'iterations' in result:
            print(f"  Итераций: {len(result['iterations'])}")
            for it in result['iterations']:
                print(f"    Итерация {it['iteration']}: "
                      f"найдено {it['issues_found']}, "
                      f"исправлено {it['issues_fixed']}")
    else:
        print(f"\nОшибка: {result.get('error', 'Неизвестная ошибка')}")
        return 1

    return 0


if __name__ == '__main__':
    exit(main())
