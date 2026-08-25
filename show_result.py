#!/usr/bin/env python3
"""Показ результата конвертации"""

from md_to_docx import RI2013Converter
from docx import Document
import os

md_path = r'd:\project\Demographics\Наборы данных\Миасский ГО\result\ОБРАЗЕЦ_Том II. Материалы по обоснованию_Миасский.md'
output_path = md_path.replace('.md', '_converted_v2.docx')
template_path = r'd:\project\Demographics\Наборы данных\Миасский ГО\result\ОБРАЗЕЦ_Том II. Материалы по обоснованию_Миасский.docx'

converter = RI2013Converter(template_path=template_path)
converter.convert(md_path, output_path)

doc = Document(output_path)

print("=" * 80)
print("РЕЗУЛЬТАТ КОНВЕРТАЦИИ MD -> DOCX")
print("=" * 80)
print(f"Файл: {output_path}")
print(f"Параграфов: {len(doc.paragraphs)}, Таблиц: {len(doc.tables)}")
print()

# Показ структуры документа
print("СТРУКТУРА ДОКУМЕНТА:")
print("-" * 80)
for i, p in enumerate(doc.paragraphs[:80]):
    text = p.text.strip()
    if text:
        style = p.style.name
        # Сокращаем длинные строки
        display_text = text[:80] + "..." if len(text) > 80 else text
        print(f"[{i:3d}] ({style:20s}) {display_text}")

print()
print("=" * 80)
print("ТАБЛИЦЫ")
print("=" * 80)
for i, table in enumerate(doc.tables[:5]):
    print(f"\nТаблица {i+1} ({len(table.rows)}x{len(table.columns)}):")
    for j, row in enumerate(table.rows[:3]):
        cells = [c.text[:15] for c in row.cells]
        print(f"  Строка {j}: {cells}")

print()
print("=" * 80)
print("ИСПОЛЬЗОВАННЫЕ СТИЛИ")
print("=" * 80)
styles_used = {}
for p in doc.paragraphs:
    if p.text.strip():
        s = p.style.name
        styles_used[s] = styles_used.get(s, 0) + 1
for s, c in sorted(styles_used.items(), key=lambda x: -x[1]):
    print(f"  {s:30s}: {c}")

# Сохраняем путь для открытия
print(f"\nФайл сохранён: {output_path}")
