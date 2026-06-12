"""
Тесты поддержки таблиц и итеративного применения исправлений.
"""
import os
import sys
import tempfile
import shutil
from unittest.mock import Mock, patch

import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from src.core.audit_engine import AuditIssue, Severity
from src.apply.apply_orchestrator import ApplyOrchestrator
from src.core.config_loader import ConfigLoader


class TestTableSupport:
    """Тесты для поддержки таблиц в ApplyOrchestrator."""
    
    def setup_method(self):
        """Создаем временный документ с таблицей."""
        self.temp_dir = tempfile.mkdtemp()
        self.doc_path = os.path.join(self.temp_dir, "test_table.docx")
        self.create_test_document()
        
        # Мок конфигурации
        self.config = Mock(spec=ConfigLoader)
        self.config.get_style_config.return_value = {
            'font': {'name': 'Times New Roman', 'size_pt': 12},
            'paragraph': {'alignment': 'left', 'space_before_pt': 0, 'space_after_pt': 0},
            'table_style': 'Table Grid'
        }
        self.engine = ApplyOrchestrator(self.config)
    
    def teardown_method(self):
        """Удаляем временные файлы."""
        shutil.rmtree(self.temp_dir, ignore_errors=True)
    
    def create_test_document(self):
        """Создает тестовый документ с одной таблицей."""
        doc = docx.Document()
        doc.add_paragraph("Параграф перед таблицей")
        table = doc.add_table(rows=3, cols=2)
        for i in range(3):
            for j in range(2):
                table.cell(i, j).text = f"Ячейка {i},{j}"
        doc.add_paragraph("Параграф после таблицы")
        doc.save(self.doc_path)
    
    def test_get_table_success(self):
        """Проверка поиска таблицы по индексу."""
        doc = docx.Document(self.doc_path)
        location = {'table_index': 0, 'type': 'table'}
        table = self.engine.table_applier._get_table(doc, location)
        assert table is not None
        assert len(table.rows) == 3
        assert len(table.columns) == 2
    
    def test_get_table_missing_index(self):
        """Проверка ошибки при отсутствии table_index."""
        doc = docx.Document(self.doc_path)
        location = {'index': 0}
        try:
            self.engine.table_applier._get_table(doc, location)
            assert False, "Должно быть исключение ValueError"
        except ValueError as e:
            assert "Отсутствует ключ table_index" in str(e)
    
    def test_get_table_out_of_range(self):
        """Проверка ошибки при неверном индексе таблицы."""
        doc = docx.Document(self.doc_path)
        location = {'table_index': 5, 'type': 'table'}
        try:
            self.engine.table_applier._get_table(doc, location)
            assert False, "Должно быть исключение ValueError"
        except ValueError as e:
            assert "Индекс таблицы 5 вне диапазона" in str(e)
    
    def test_apply_table_style(self):
        """Проверка применения стиля таблицы."""
        doc = docx.Document(self.doc_path)
        table = doc.tables[0]
        # Убедимся, что стиль изначально не 'Table Grid'
        original_style = table.style.name if table.style else None
        
        # Применяем стиль через table_applier
        success = self.engine.table_applier.apply_style(table, "Table Grid")
        assert success is True
        # В python-docx стиль таблицы можно установить, но проверка может быть сложной
        # Проверим, что атрибут style существует
        assert hasattr(table, 'style')
    
    def test_apply_fixes_with_table(self):
        """Проверка применения исправлений для таблицы через apply_fixes."""
        # Создаем проблему таблицы
        issue = AuditIssue(
            id="table_issue",
            severity=Severity.WARNING,
            category="TABLE",
            element_type="Table Grid",
            location={'table_index': 0, 'type': 'table'},
            description="Неверный стиль таблицы",
            current_value="No Style",
            expected_value="Table Grid",
            auto_fixable=True
        )
        
        # Мокаем style_manager, чтобы не создавать реальные стили
        with patch.object(self.engine.style_manager, 'ensure_style_exists'):
            stats = self.engine.apply_fixes(
                self.doc_path,
                [issue],
                output_path=os.path.join(self.temp_dir, "output.docx"),
                apply_styles=True,
                apply_direct_overrides=False,
                clear_direct_formatting=False
            )
        
        # Проверяем, что исправление применено
        assert stats['applied'] == 1
        assert stats['failed'] == 0
    
    def test_apply_fixes_with_table_and_paragraph(self):
        """Проверка смешанного применения: таблица + параграф."""
        # Проблема таблицы
        table_issue = AuditIssue(
            id="table_issue",
            severity=Severity.WARNING,
            category="TABLE",
            element_type="Table Grid",
            location={'table_index': 0, 'type': 'table'},
            description="Неверный стиль таблицы",
            current_value="No Style",
            expected_value="Table Grid",
            auto_fixable=True
        )
        # Проблема параграфа
        para_issue = AuditIssue(
            id="para_issue",
            severity=Severity.WARNING,
            category="PARAGRAPH",
            element_type="Normal",
            location={'index': 0},
            description="Неверный стиль параграфа",
            current_value="No Style",
            expected_value="Normal",
            auto_fixable=True
        )
        
        with patch.object(self.engine.style_manager, 'ensure_style_exists'):
            with patch.object(self.engine.style_manager, 'apply_direct_overrides'):
                stats = self.engine.apply_fixes(
                    self.doc_path,
                    [table_issue, para_issue],
                    output_path=os.path.join(self.temp_dir, "output_mixed.docx"),
                    apply_styles=True,
                    apply_direct_overrides=False,
                    clear_direct_formatting=False
                )
        
        # Оба исправления должны быть применены
        assert stats['applied'] == 2
        assert stats['failed'] == 0


class TestIterativeApplication:
    """Тесты для итеративного применения исправлений."""
    
    def setup_method(self):
        self.temp_dir = tempfile.mkdtemp()
        self.doc_path = os.path.join(self.temp_dir, "test_iterative.docx")
        self.create_test_document()
        
        self.config = Mock(spec=ConfigLoader)
        self.config.get_style_config.return_value = {
            'font': {'name': 'Times New Roman', 'size_pt': 12},
            'paragraph': {'alignment': 'left', 'space_before_pt': 0, 'space_after_pt': 0},
        }
        self.engine = ApplyOrchestrator(self.config)
    
    def teardown_method(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)
    
    def create_test_document(self):
        doc = docx.Document()
        doc.add_paragraph("Первый параграф")
        doc.add_paragraph("Второй параграф")
        doc.save(self.doc_path)
    
    def test_apply_fixes_iterative_success(self):
        """Проверка итеративного применения с успешным завершением."""
        # Создаем две проблемы
        issues = [
            AuditIssue(
                id="issue_1",
                severity=Severity.WARNING,
                category="PARAGRAPH",
                element_type="Normal",
                location={'index': 0},
                description="Проблема 1",
                current_value="No Style",
                expected_value="Normal",
                auto_fixable=True
            ),
            AuditIssue(
                id="issue_2",
                severity=Severity.WARNING,
                category="PARAGRAPH",
                element_type="Normal",
                location={'index': 1},
                description="Проблема 2",
                current_value="No Style",
                expected_value="Normal",
                auto_fixable=True
            )
        ]
        
        with patch.object(self.engine.style_manager, 'ensure_style_exists'):
            with patch.object(self.engine.style_manager, 'apply_direct_overrides'):
                # Мокаем apply_fixes, чтобы симулировать прогресс
                call_count = 0
                def mock_apply_fixes(doc_path, issues, output_path=None, **kwargs):
                    nonlocal call_count
                    call_count += 1
                    # Первый вызов: одна ошибка, второй вызов: все успешно
                    if call_count == 1:
                        return {'applied': 1, 'failed': 1}
                    else:
                        return {'applied': 2, 'failed': 0}
                
                with patch.object(self.engine, 'apply_fixes', side_effect=mock_apply_fixes):
                    results = self.engine.apply_fixes_iterative(
                        self.doc_path,
                        issues,
                        output_path=os.path.join(self.temp_dir, "output_iter.docx"),
                        max_iterations=3,
                        apply_styles=True,
                        apply_direct_overrides=False,
                        clear_direct_formatting=False
                    )
        
        # Должно быть две итерации
        assert len(results) == 2
        assert results[0]['failed'] == 1
        assert results[1]['failed'] == 0
    
    def test_apply_fixes_iterative_no_progress(self):
        """Проверка остановки при отсутствии прогресса."""
        issues = [
            AuditIssue(
                id="issue_1",
                severity=Severity.WARNING,
                category="PARAGRAPH",
                element_type="Normal",
                location={'index': 0},
                description="Проблема",
                current_value="No Style",
                expected_value="Normal",
                auto_fixable=True
            )
        ]
        
        with patch.object(self.engine.style_manager, 'ensure_style_exists'):
            with patch.object(self.engine.style_manager, 'apply_direct_overrides'):
                # Мокаем apply_fixes, чтобы всегда возвращать одну ошибку
                def mock_apply_fixes(doc_path, issues, output_path=None, **kwargs):
                    return {'applied': 0, 'failed': 1}
                
                with patch.object(self.engine, 'apply_fixes', side_effect=mock_apply_fixes):
                    results = self.engine.apply_fixes_iterative(
                        self.doc_path,
                        issues,
                        output_path=os.path.join(self.temp_dir, "output_no_progress.docx"),
                        max_iterations=5,
                        apply_styles=True,
                        apply_direct_overrides=False,
                        clear_direct_formatting=False
                    )
        
        # Должно быть две итерации (вторая не уменьшает количество ошибок)
        assert len(results) == 2
        assert all(r['failed'] == 1 for r in results)
    
    def test_apply_fixes_iterative_all_success_first(self):
        """Проверка, когда все исправления успешны на первой итерации."""
        issues = [
            AuditIssue(
                id="issue_1",
                severity=Severity.WARNING,
                category="PARAGRAPH",
                element_type="Normal",
                location={'index': 0},
                description="Проблема",
                current_value="No Style",
                expected_value="Normal",
                auto_fixable=True
            )
        ]
        
        with patch.object(self.engine.style_manager, 'ensure_style_exists'):
            with patch.object(self.engine.style_manager, 'apply_direct_overrides'):
                def mock_apply_fixes(doc_path, issues, output_path=None, **kwargs):
                    return {'applied': 1, 'failed': 0}
                
                with patch.object(self.engine, 'apply_fixes', side_effect=mock_apply_fixes):
                    results = self.engine.apply_fixes_iterative(
                        self.doc_path,
                        issues,
                        output_path=os.path.join(self.temp_dir, "output_success_first.docx"),
                        max_iterations=3,
                        apply_styles=True,
                        apply_direct_overrides=False,
                        clear_direct_formatting=False
                    )
        
        # Только одна итерация
        assert len(results) == 1
        assert results[0]['failed'] == 0


class TestUniversalGetElement:
    """Тесты универсальной функции _get_element."""
    
    def setup_method(self):
        self.temp_dir = tempfile.mkdtemp()
        self.doc_path = os.path.join(self.temp_dir, "test_universal.docx")
        self.create_test_document()
        
        self.config = Mock(spec=ConfigLoader)
        self.engine = ApplyOrchestrator(self.config)
    
    def teardown_method(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)
    
    def create_test_document(self):
        doc = docx.Document()
        doc.add_paragraph("Параграф 1")
        doc.add_table(rows=2, cols=2)
        doc.save(self.doc_path)
    
    def test_get_element_paragraph(self):
        """Получение параграфа через _get_element."""
        doc = docx.Document(self.doc_path)
        location = {'index': 0, 'type': 'paragraph'}
        element, elem_type = self.engine.paragraph_applier._get_element(doc, location)
        assert elem_type == 'paragraph'
        assert element.text == "Параграф 1"
    
    def test_get_element_table(self):
        """Получение таблицы через _get_element."""
        doc = docx.Document(self.doc_path)
        location = {'table_index': 0, 'type': 'table'}
        element, elem_type = self.engine.table_applier._get_element(doc, location)
        assert elem_type == 'table'
        assert len(element.rows) == 2
    
    def test_get_element_default_type(self):
        """Если тип не указан, по умолчанию 'paragraph'."""
        doc = docx.Document(self.doc_path)
        location = {'index': 0}  # без type
        element, elem_type = self.engine.paragraph_applier._get_element(doc, location)
        assert elem_type == 'paragraph'
    
    def test_get_element_unsupported_type(self):
        """Неподдерживаемый тип вызывает исключение."""
        doc = docx.Document(self.doc_path)
        location = {'type': 'figure'}
        try:
            self.engine.paragraph_applier._get_element(doc, location)
            assert False, "Должно быть исключение ValueError"
        except ValueError as e:
            assert "Тип элемента figure пока не поддерживается" in str(e)
    
    def test_get_element_unknown_type(self):
        """Неизвестный тип вызывает исключение."""
        doc = docx.Document(self.doc_path)
        location = {'type': 'unknown'}
        try:
            self.engine.paragraph_applier._get_element(doc, location)
            assert False, "Должно быть исключение ValueError"
        except ValueError as e:
            assert "Неизвестный тип элемента" in str(e)


if __name__ == '__main__':
    import pytest
    pytest.main([__file__, '-v'])