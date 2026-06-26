#!/usr/bin/env python3
"""
Snapshot-тесты для XML-структуры DOCX.
Сравнивает сгенерированные документы с эталонными снимками.
"""

import os
import sys
import tempfile
import shutil
import pytest
from docx import Document

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(__file__))))

from src.core.snapshot_testing import DocxSnapshot, get_docx_styles, get_docx_paragraph_count


SNAPSHOTS_DIR = os.path.join(os.path.dirname(__file__), "..", "snapshots")
UPDATE_SNAPSHOTS = os.environ.get("UPDATE_SNAPSHOTS", "0") == "1"


@pytest.fixture
def snapshot():
    return DocxSnapshot(SNAPSHOTS_DIR, update_mode=UPDATE_SNAPSHOTS)


@pytest.fixture
def temp_dir():
    d = tempfile.mkdtemp()
    yield d
    shutil.rmtree(d, ignore_errors=True)


class TestDocxStructure:
    """Тесты структуры DOCX после форматирования."""

    def test_empty_document_structure(self, snapshot, temp_dir):
        """Проверяет XML-структуру пустого документа."""
        doc = Document()
        doc_path = os.path.join(temp_dir, "empty.docx")
        doc.save(doc_path)

        snapshot.assert_matches("empty_document", doc_path)

    def test_paragraph_with_style(self, snapshot, temp_dir):
        """Проверяет структуру параграфа со стилем."""
        doc = Document()
        p = doc.add_paragraph("Тестовый текст")
        p.style = doc.styles['Normal']
        doc_path = os.path.join(temp_dir, "styled_para.docx")
        doc.save(doc_path)

        snapshot.assert_matches("paragraph_with_style", doc_path)

    def test_heading_styles(self, snapshot, temp_dir):
        """Проверяет структуру заголовков разных уровней."""
        doc = Document()
        doc.add_heading("Заголовок 1", level=1)
        doc.add_heading("Заголовок 2", level=2)
        doc.add_heading("Заголовок 3", level=3)
        doc_path = os.path.join(temp_dir, "headings.docx")
        doc.save(doc_path)

        snapshot.assert_matches("heading_styles", doc_path)

    def test_table_structure(self, snapshot, temp_dir):
        """Проверяет структуру таблицы."""
        doc = Document()
        table = doc.add_table(rows=3, cols=3)
        for i, row in enumerate(table.rows):
            for j, cell in enumerate(row.cells):
                cell.text = f"Ячейка {i},{j}"
        doc_path = os.path.join(temp_dir, "table.docx")
        doc.save(doc_path)

        snapshot.assert_matches("table_structure", doc_path)

    def test_list_items(self, snapshot, temp_dir):
        """Проверяет структуру списков."""
        doc = Document()
        doc.add_paragraph("Элемент 1", style='List Bullet')
        doc.add_paragraph("Элемент 2", style='List Bullet')
        doc.add_paragraph("Элемент 3", style='List Number')
        doc_path = os.path.join(temp_dir, "lists.docx")
        doc.save(doc_path)

        snapshot.assert_matches("list_items", doc_path)


class TestFormattedDocument:
    """Тесты XML-структуры после форматирования по ГОСТ."""

    def test_config_formatting(self, snapshot, temp_dir):
        """Проверяет структуру после форматирования через конфиг."""
        from src.core.config_loader import ConfigLoader
        from src.apply.apply_orchestrator import ApplyOrchestrator

        config = ConfigLoader("configs/active/config.yaml")
        config.load()

        doc = Document()
        doc.add_paragraph("Текст для форматирования")
        input_path = os.path.join(temp_dir, "input.docx")
        doc.save(input_path)

        output_path = os.path.join(temp_dir, "output.docx")
        orchestrator = ApplyOrchestrator(config)

        from src.core.audit_engine import AuditEngine
        audit = AuditEngine(config)
        issues = audit.scan_document(input_path)
        fixable = [i for i in issues if i.auto_fixable]

        if fixable:
            orchestrator.apply_fixes(input_path, fixable, output_path)

        if os.path.exists(output_path):
            snapshot.assert_matches("config_formatted", output_path)

    def test_template_formatting(self, snapshot, temp_dir):
        """Проверяет структуру после форматирования через шаблон."""
        template_path = "НТД 01-2013 Текстовая часть_4ред.docx"
        if not os.path.exists(template_path):
            pytest.skip("Шаблон не найден")

        from src.apply.style_applier import StyleApplier

        doc = Document()
        doc.add_paragraph("Текст для шаблона")
        input_path = os.path.join(temp_dir, "input.docx")
        doc.save(input_path)

        output_path = os.path.join(temp_dir, "output.docx")
        applier = StyleApplier(template_path)
        applier.apply(input_path, output_path)

        snapshot.assert_matches("template_formatted", output_path)


class TestDocxUtils:
    """Тесты утилитарных функций."""

    def test_get_docx_styles(self, temp_dir):
        """Проверяет извлечение стилей из DOCX."""
        doc = Document()
        doc.add_paragraph("Текст")
        doc_path = os.path.join(temp_dir, "test.docx")
        doc.save(doc_path)

        styles = get_docx_styles(doc_path)
        assert len(styles) > 0
        assert any(s["name"] == "Normal" for s in styles)

    def test_get_docx_paragraph_count(self, temp_dir):
        """Проверяет подсчёт параграфов."""
        doc = Document()
        doc.add_paragraph("Параграф 1")
        doc.add_paragraph("Параграф 2")
        doc.add_paragraph("Параграф 3")
        doc_path = os.path.join(temp_dir, "test.docx")
        doc.save(doc_path)

        count = get_docx_paragraph_count(doc_path)
        assert count == 3
