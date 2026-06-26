#!/usr/bin/env python3
"""Расширенные тесты для docx_utils."""

import os
import tempfile
import shutil
import pytest
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from src.core.docx_utils import (
    get_xml_attr_twips, get_paragraph_indent_twips,
    get_paragraph_spacing_twips, get_paragraph_alignment,
    get_font_properties, normalize_color, open_document,
    ALIGNMENT_MAP, cm_to_twips, pt_to_twips, twips_to_cm, twips_to_pt
)
from docx.shared import RGBColor


class TestGetXmlAttrTwips:
    """Тесты get_xml_attr_twips."""

    def test_valid_attr(self):
        elem = OxmlElement('w:ind')
        elem.set(qn('w:firstLine'), '720')
        result = get_xml_attr_twips(elem, 'w:firstLine')
        assert result == 720

    def test_missing_attr(self):
        elem = OxmlElement('w:ind')
        result = get_xml_attr_twips(elem, 'w:firstLine')
        assert result is None

    def test_none_element(self):
        result = get_xml_attr_twips(None, 'w:firstLine')
        assert result is None

    def test_invalid_value(self):
        elem = OxmlElement('w:ind')
        elem.set(qn('w:firstLine'), 'abc')
        result = get_xml_attr_twips(elem, 'w:firstLine')
        assert result is None


class TestGetParagraphIndentTwips:
    """Тесты get_paragraph_indent_twips."""

    def test_with_first_line(self):
        doc = Document()
        para = doc.add_paragraph('Test')
        pPr = para._element.get_or_add_pPr()
        ind = OxmlElement('w:ind')
        ind.set(qn('w:firstLine'), '720')
        pPr.append(ind)

        result = get_paragraph_indent_twips(para)
        assert result['first_line'] == 720

    def test_without_indent(self):
        doc = Document()
        para = doc.add_paragraph('Test')
        result = get_paragraph_indent_twips(para)
        assert result['first_line'] == 0


class TestGetParagraphSpacingTwips:
    """Тесты get_paragraph_spacing_twips."""

    def test_with_spacing(self):
        doc = Document()
        para = doc.add_paragraph('Test')
        pPr = para._element.get_or_add_pPr()
        spacing = OxmlElement('w:spacing')
        spacing.set(qn('w:before'), '240')
        spacing.set(qn('w:after'), '120')
        pPr.append(spacing)

        result = get_paragraph_spacing_twips(para)
        assert result['before'] == 240
        assert result['after'] == 120

    def test_without_spacing(self):
        doc = Document()
        para = doc.add_paragraph('Test')
        result = get_paragraph_spacing_twips(para)
        assert result['before'] == 0
        assert result['after'] == 0


class TestGetParagraphAlignment:
    """Тесты get_paragraph_alignment."""

    def test_justify(self):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        doc = Document()
        para = doc.add_paragraph('Test')
        para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        result = get_paragraph_alignment(para)
        assert result == 'justify'

    def test_center(self):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        doc = Document()
        para = doc.add_paragraph('Test')
        para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        result = get_paragraph_alignment(para)
        assert result == 'center'

    def test_left_default(self):
        doc = Document()
        para = doc.add_paragraph('Test')
        result = get_paragraph_alignment(para)
        assert result == 'left'


class TestGetFontProperties:
    """Тесты get_font_properties."""

    def test_with_font(self):
        doc = Document()
        para = doc.add_paragraph('Test')
        run = para.add_run('Run')
        run.font.name = 'Arial'
        run.font.size = Pt(14)
        run.font.bold = True

        props = get_font_properties(run)
        assert props['name'] == 'Arial'
        assert props['bold'] is True

    def test_empty_run(self):
        doc = Document()
        para = doc.add_paragraph('')
        if para.runs:
            run = para.runs[0]
        else:
            run = para.add_run('')

        props = get_font_properties(run)
        assert props['name'] is None


class TestNormalizeColor:
    """Тесты normalize_color."""

    def test_hex_6(self):
        result = normalize_color('#FF0000')
        assert result == RGBColor(255, 0, 0)

    def test_hex_3(self):
        result = normalize_color('#F00')
        assert result == RGBColor(255, 0, 0)

    def test_named_color(self):
        result = normalize_color('red')
        assert result == RGBColor(255, 0, 0)

    def test_none(self):
        result = normalize_color(None)
        assert result is None

    def test_invalid(self):
        result = normalize_color('invalid')
        assert result is None


class TestOpenDocument:
    """Тесты open_document."""

    def setup_method(self):
        self.temp_dir = tempfile.mkdtemp()

    def teardown_method(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_open_existing(self):
        path = os.path.join(self.temp_dir, 'test.docx')
        doc = Document()
        doc.add_paragraph('Test')
        doc.save(path)

        result = open_document(path)
        assert result is not None

    def test_open_nonexistent(self):
        from src.core.formatter_errors import FormatterError
        with pytest.raises(FormatterError):
            open_document('/nonexistent/file.docx')


class TestAlignmentMap:
    """Тесты ALIGNMENT_MAP."""

    def test_map_values(self):
        assert ALIGNMENT_MAP['both'] == 'justify'
        assert ALIGNMENT_MAP['left'] == 'left'
        assert ALIGNMENT_MAP['center'] == 'center'
        assert ALIGNMENT_MAP['right'] == 'right'
        assert ALIGNMENT_MAP[None] == 'left'


class TestConversionFunctions:
    """Тесты функций конвертации."""

    def test_cm_to_twips(self):
        assert cm_to_twips(1) == 567

    def test_pt_to_twips(self):
        assert pt_to_twips(1) == 20

    def test_twips_to_cm(self):
        assert twips_to_cm(567) == 1.0

    def test_twips_to_pt(self):
        assert twips_to_pt(20) == 1.0
