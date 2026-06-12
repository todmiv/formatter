"""
Тесты для улучшенного модуля применения (apply_engine.py) с новыми функциями:
- Селективное применение исправлений (только стили)
- Опции управления прямым форматированием
- Оптимизация производительности
"""
import sys
import os
import tempfile
import shutil
import time
sys.path.insert(0, '.')

import docx
from src.core.config_loader import ConfigLoader
from src.core.audit_engine import AuditEngine, AuditIssue, Severity
from src.apply.apply_orchestrator import ApplyOrchestrator

def create_test_document(path):
    """Создаёт тестовый документ с известными проблемами."""
    doc = docx.Document()
    # Параграф с неправильным стилем
    p1 = doc.add_paragraph("Заголовок 1")
    p1.style = doc.styles['Normal']
    # Параграф с прямым форматированием
    p2 = doc.add_paragraph("Текст с жирным шрифтом")
    run = p2.add_run("жирный")
    run.font.bold = True
    run.font.size = docx.shared.Pt(20)
    # Параграф с отступом
    p3 = doc.add_paragraph("Текст с отступом")
    p3.paragraph_format.left_indent = docx.shared.Cm(2)
    doc.save(path)
    return path

def test_apply_styles_only():
    """Тест применения только стилей без прямого форматирования."""
    print("Тест: apply_styles_only")
    with tempfile.TemporaryDirectory() as tmpdir:
        doc_path = os.path.join(tmpdir, "test.docx")
        create_test_document(doc_path)
        
        loader = ConfigLoader("config.yaml")
        loader.load()
        
        # Создаём искусственные проблемы
        issues = [
            AuditIssue(
                id="test_1",
                severity=Severity.WARNING,
                category="STYLE",
                element_type="Heading 1",
                location={"index": 0},
                description="Неправильный стиль",
                current_value="Normal",
                expected_value="Heading 1",
                auto_fixable=True
            ),
            AuditIssue(
                id="test_2",
                severity=Severity.WARNING,
                category="FONT",
                element_type="Normal",
                location={"index": 1},
                description="Неправильный размер шрифта",
                current_value="20pt",
                expected_value="14pt",
                auto_fixable=True
            )
        ]
        
        engine = ApplyOrchestrator(loader)
        output_path = os.path.join(tmpdir, "fixed.docx")
        
        # Применяем только стили
        stats = engine.apply_fixes(doc_path, issues, output_path,
                                   apply_styles=True,
                                   apply_direct_overrides=False,
                                   clear_direct_formatting=False)
        
        assert stats['applied'] == 2
        assert stats['failed'] == 0
        
        # Проверяем, что документ создан
        assert os.path.exists(output_path)
        
        # Загружаем исправленный документ и проверяем стили
        doc = docx.Document(output_path)
        # Первый параграф должен иметь стиль Heading 1
        assert doc.paragraphs[0].style.name == "Heading 1"
        # Второй параграф должен иметь стиль Normal (но прямое форматирование осталось)
        assert doc.paragraphs[1].style.name == "Normal"
        # Проверяем, что прямое форматирование не было изменено (жирный остался)
        # В python-docx проверить сложно, но хотя бы убедимся, что документ открывается
        print("  [OK] apply_styles_only пройден")

def test_apply_with_direct_overrides():
    """Тест применения стилей с прямым форматированием."""
    print("Тест: apply_with_direct_overrides")
    with tempfile.TemporaryDirectory() as tmpdir:
        doc_path = os.path.join(tmpdir, "test.docx")
        create_test_document(doc_path)
        
        loader = ConfigLoader("config.yaml")
        loader.load()
        
        issues = [
            AuditIssue(
                id="test_1",
                severity=Severity.WARNING,
                category="FONT",
                element_type="Normal",
                location={"index": 1},
                description="Неправильный размер шрифта",
                current_value="20pt",
                expected_value="14pt",
                auto_fixable=True
            )
        ]
        
        engine = ApplyOrchestrator(loader)
        output_path = os.path.join(tmpdir, "fixed.docx")
        
        # Применяем с прямым форматированием
        stats = engine.apply_fixes(doc_path, issues, output_path,
                                   apply_styles=True,
                                   apply_direct_overrides=True,
                                   clear_direct_formatting=False)
        
        assert stats['applied'] == 1
        assert stats['failed'] == 0
        assert os.path.exists(output_path)
        print("  [OK] apply_with_direct_overrides пройден")

def test_clear_direct_formatting():
    """Тест очистки прямого форматирования."""
    print("Тест: clear_direct_formatting")
    with tempfile.TemporaryDirectory() as tmpdir:
        doc_path = os.path.join(tmpdir, "test.docx")
        create_test_document(doc_path)
        
        loader = ConfigLoader("config.yaml")
        loader.load()
        
        issues = [
            AuditIssue(
                id="test_1",
                severity=Severity.WARNING,
                category="FONT",
                element_type="Normal",
                location={"index": 1},
                description="Неправильный размер шрифта",
                current_value="20pt",
                expected_value="14pt",
                auto_fixable=True
            )
        ]
        
        engine = ApplyOrchestrator(loader)
        output_path = os.path.join(tmpdir, "fixed.docx")
        
        # Применяем с очисткой прямого форматирования
        stats = engine.apply_fixes(doc_path, issues, output_path,
                                   apply_styles=True,
                                   apply_direct_overrides=True,
                                   clear_direct_formatting=True)
        
        assert stats['applied'] == 1
        assert stats['failed'] == 0
        assert os.path.exists(output_path)
        print("  [OK] clear_direct_formatting пройден")

def test_apply_styles_only_method():
    """Тест метода apply_styles_only."""
    print("Тест: apply_styles_only метод")
    with tempfile.TemporaryDirectory() as tmpdir:
        doc_path = os.path.join(tmpdir, "test.docx")
        create_test_document(doc_path)
        
        loader = ConfigLoader("config.yaml")
        loader.load()
        
        issues = [
            AuditIssue(
                id="test_1",
                severity=Severity.WARNING,
                category="STYLE",
                element_type="Heading 1",
                location={"index": 0},
                description="Неправильный стиль",
                current_value="Normal",
                expected_value="Heading 1",
                auto_fixable=True
            )
        ]
        
        engine = ApplyOrchestrator(loader)
        output_path = os.path.join(tmpdir, "fixed.docx")
        
        stats = engine.apply_styles_only(doc_path, issues, output_path)
        
        assert stats['applied'] == 1
        assert stats['failed'] == 0
        assert os.path.exists(output_path)
        print("  [OK] apply_styles_only метод пройден")

def test_performance_optimization():
    """Тест оптимизации производительности (группировка и кэширование)."""
    print("Тест: производительность")
    with tempfile.TemporaryDirectory() as tmpdir:
        doc_path = os.path.join(tmpdir, "test.docx")
        # Создаём документ с большим количеством параграфов
        doc = docx.Document()
        for i in range(100):
            p = doc.add_paragraph(f"Параграф {i}")
            p.style = doc.styles['Normal']
        doc.save(doc_path)
        
        loader = ConfigLoader("config.yaml")
        loader.load()
        
        # Создаём 100 проблем
        issues = []
        for i in range(100):
            issues.append(
                AuditIssue(
                    id=f"test_{i}",
                    severity=Severity.WARNING,
                    category="STYLE",
                    element_type="Heading 1",
                    location={"index": i},
                    description="Неправильный стиль",
                    current_value="Normal",
                    expected_value="Heading 1",
                    auto_fixable=True
                )
            )
        
        engine = ApplyOrchestrator(loader)
        output_path = os.path.join(tmpdir, "fixed.docx")
        
        start_time = time.time()
        stats = engine.apply_fixes(doc_path, issues, output_path,
                                   apply_styles=True,
                                   apply_direct_overrides=False,
                                   clear_direct_formatting=False)
        end_time = time.time()
        
        duration = end_time - start_time
        print(f"  Время применения 100 исправлений: {duration:.2f} сек")
        assert stats['applied'] == 100
        assert stats['failed'] == 0
        assert duration < 10.0  # Должно быть достаточно быстро
        print("  [OK] производительность пройдена")

def test_main_controller_integration():
    """Тест интеграции с MainController."""
    print("Тест: интеграция с MainController")
    from src.core.main_controller import MainController
    
    with tempfile.TemporaryDirectory() as tmpdir:
        doc_path = os.path.join(tmpdir, "test.docx")
        create_test_document(doc_path)
        
        controller = MainController("config.yaml")
        controller.set_document(doc_path)
        
        # Создаём искусственные проблемы
        issues = [
            AuditIssue(
                id="test_1",
                severity=Severity.WARNING,
                category="STYLE",
                element_type="Heading 1",
                location={"index": 0},
                description="Неправильный стиль",
                current_value="Normal",
                expected_value="Heading 1",
                auto_fixable=True
            )
        ]
        controller.current_issues = issues
        
        # Тестируем apply_styles_only
        stats = controller.apply_styles_only(["test_1"])
        assert stats['applied'] == 1
        assert stats['failed'] == 0
        
        # Тестируем apply_fixes_with_options
        issues2 = [
            AuditIssue(
                id="test_2",
                severity=Severity.WARNING,
                category="FONT",
                element_type="Normal",
                location={"index": 1},
                description="Неправильный размер",
                current_value="20pt",
                expected_value="14pt",
                auto_fixable=True
            )
        ]
        controller.current_issues = issues2
        stats = controller.apply_fixes_with_options(["test_2"], 
                                                    apply_styles=True,
                                                    apply_direct_overrides=False,
                                                    clear_direct_formatting=True)
        assert stats['applied'] == 1
        assert stats['failed'] == 0
        
        print("  [OK] интеграция с MainController пройдена")

def run_all_tests():
    """Запускает все тесты."""
    print("=== Запуск тестов улучшенного модуля применения ===")
    try:
        test_apply_styles_only()
        test_apply_with_direct_overrides()
        test_clear_direct_formatting()
        test_apply_styles_only_method()
        test_performance_optimization()
        test_main_controller_integration()
        print("\n=== Все тесты пройдены успешно ===")
        return True
    except Exception as e:
        print(f"\n=== Тест провален: {e} ===")
        import traceback
        traceback.print_exc()
        return False

if __name__ == "__main__":
    success = run_all_tests()
    sys.exit(0 if success else 1)