#!/usr/bin/env python3
"""Тесты для DocumentModel."""

import os
import tempfile
import shutil
import pytest
from docx import Document
from docx.shared import Pt, Cm
from src.core.document_model import DocumentModel


class TestDocumentModel:
    """Тесты для DocumentModel."""

    def setup_method(self):
        self.temp_dir = tempfile.mkdtemp()
        self.doc_path = os.path.join(self.temp_dir, 'test.docx')
        doc = Document()
        doc.add_paragraph('Para 1', style='Normal')
        doc.add_paragraph('Para 2', style='Normal')
        doc.add_paragraph('Para 3', style='Normal')
        doc.add_table(rows=2, cols=2)
        doc.save(self.doc_path)

    def teardown_method(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_init_from_path(self):
        model = DocumentModel(self.doc_path)
        assert model.path == self.doc_path
        assert model.docx_document is not None

    def test_init_new_document(self):
        model = DocumentModel()
        assert model.path is None
        assert model.docx_document is not None

    def test_init_from_doc(self):
        doc = Document()
        doc.add_paragraph('Test')
        model = DocumentModel(doc=doc)
        assert model.path is None
        assert model.docx_document is doc

    def test_paragraphs(self):
        model = DocumentModel(self.doc_path)
        paras = model.paragraphs()
        assert len(paras) == 3

    def test_paragraph_at(self):
        model = DocumentModel(self.doc_path)
        para = model.paragraph_at(0)
        assert para is not None
        assert para.text == 'Para 1'

    def test_paragraph_at_out_of_range(self):
        model = DocumentModel(self.doc_path)
        para = model.paragraph_at(999)
        assert para is None

    def test_tables(self):
        model = DocumentModel(self.doc_path)
        tables = model.tables()
        assert len(tables) == 1

    def test_sections(self):
        model = DocumentModel(self.doc_path)
        sections = model.sections()
        assert len(sections) >= 1

    def test_styles(self):
        model = DocumentModel(self.doc_path)
        styles = model.styles()
        assert styles is not None

    def test_get_paragraph_properties(self):
        model = DocumentModel(self.doc_path)
        para = model.paragraph_at(0)
        props = model.get_paragraph_properties(para)
        assert 'style' in props
        assert 'alignment' in props
        assert 'first_line_indent_twips' in props

    def test_get_run_font_properties(self):
        model = DocumentModel(self.doc_path)
        para = model.paragraph_at(0)
        run = para.add_run('Test run')
        props = model.get_run_font_properties(run)
        assert 'name' in props
        assert 'bold' in props

    def test_get_style_config(self):
        model = DocumentModel(self.doc_path)
        config = model.get_style_config('Normal')
        assert config is not None
        assert config['name'] == 'Normal'

    def test_get_style_config_nonexistent(self):
        model = DocumentModel(self.doc_path)
        config = model.get_style_config('NonExistent')
        assert config is None

    def test_apply_style(self):
        model = DocumentModel(self.doc_path)
        para = model.paragraph_at(0)
        model.apply_style(para, 'Normal')
        assert para.style.name == 'Normal'

    def test_set_paragraph_format(self):
        model = DocumentModel(self.doc_path)
        para = model.paragraph_at(0)
        model.set_paragraph_format(
            para,
            alignment='center',
            first_line_indent_cm=1.0,
            space_before_pt=6,
            line_spacing=1.5
        )
        assert para.paragraph_format.space_before is not None

    def test_set_run_font(self):
        model = DocumentModel(self.doc_path)
        para = model.paragraph_at(0)
        run = para.add_run('Test')
        model.set_run_font(
            run,
            name='Arial',
            size_pt=14,
            bold=True,
            italic=True
        )
        assert run.font.name == 'Arial'
        assert run.font.bold is True

    def test_clear_direct_formatting(self):
        model = DocumentModel(self.doc_path)
        para = model.paragraph_at(0)
        run = para.add_run('Bold text')
        run.font.bold = True
        model.clear_direct_formatting(para)
        assert para.runs[0].font.bold is None

    def test_save(self):
        model = DocumentModel(self.doc_path)
        save_path = os.path.join(self.temp_dir, 'saved.docx')
        model.save(save_path)
        assert os.path.exists(save_path)

    def test_save_overwrite(self):
        model = DocumentModel(self.doc_path)
        para = model.paragraph_at(0)
        para.text = 'Modified'
        model.save()
        model2 = DocumentModel(self.doc_path)
        assert model2.paragraph_at(0).text == 'Modified'

    def test_save_no_path_raises(self):
        model = DocumentModel()
        with pytest.raises(ValueError):
            model.save()

    def test_estimate_page_number(self):
        model = DocumentModel(self.doc_path)
        page = model.estimate_page_number(50, lines_per_page=25)
        assert page == 3

    def test_context_manager(self):
        with DocumentModel(self.doc_path) as model:
            assert model.docx_document is not None
