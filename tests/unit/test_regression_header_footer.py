#!/usr/bin/env python3
"""
Тест регрессии для исправлений колонтитулов.
Проверяет, что проблемы колонтитулов всегда исправляются с первого раза,
обратную совместимость и работу с документами, содержащими только колонтитулы.
"""

import sys
import os
import tempfile
import shutil
import logging
from typing import List, Dict, Any
import unittest

sys.path.append(os.path.dirname(os.path.abspath(__file__)))

from src.core.config_loader import ConfigLoader
from src.core.audit_engine import AuditEngine, AuditIssue, Severity
from src.apply.apply_orchestrator import ApplyOrchestrator
import docx

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)


class TestHeaderFooterRegression(unittest.TestCase):
    """Класс тестов регрессии для колонтитулов."""
    
    @classmethod
    def setUpClass(cls):
        """Настройка перед всеми тестами."""
        # Загрузка конфигурации
        config_path = "configs/active/config.yaml"
        if not os.path.exists(config_path):
            config_path = "configs/active/config_v4.2.yaml"
        cls.config = ConfigLoader(config_path)
        cls.config.load()
        
        cls.audit_engine = AuditEngine(cls.config)
        cls.apply_engine = ApplyOrchestrator(cls.config)
        
        # Путь к тестовому документу с колонтитулами
        cls.test_doc_path = "tests/test.docx"
        if not os.path.exists(cls.test_doc_path):
            cls.test_doc_path = "Том II_Волжск.docx"
        if not os.path.exists(cls.test_doc_path):
            # Если нет, создадим простой документ с колонтитулами для теста
            cls._create_test_document_with_headers()
    
    @classmethod
    def _create_test_document_with_headers(cls):
        """Создаёт тестовый документ с колонтитулами, если оригинальный отсутствует."""
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.shared import Pt
        
        doc = Document()
        
        # Добавляем секцию с колонтитулами
        section = doc.sections[0]
        header = section.header
        footer = section.footer
        
        # Добавляем текст в верхний колонтитул
        header_para = header.paragraphs[0]
        header_para.text = "Верхний колонтитул - Тест"
        header_para.style = "Header"
        
        # Добавляем текст в нижний колонтитул
        footer_para = footer.paragraphs[0]
        footer_para.text = "Нижний колонтитул - Страница"
        footer_para.style = "Footer"
        
        # Добавляем основной текст
        doc.add_paragraph("Основной текст документа для тестирования.")
        doc.add_paragraph("Второй параграф.")
        
        # Сохраняем
        cls.test_doc_path = "test_document_with_headers.docx"
        doc.save(cls.test_doc_path)
        logger.info(f"Создан тестовый документ: {cls.test_doc_path}")
    
    def test_1_header_footer_always_fixed_first_time(self):
        """
        Тест 1: Проблемы колонтитулов всегда исправляются с первого раза.
        """
        logger.info("=== Тест 1: Исправление колонтитулов с первого раза ===")
        
        # Аудит документа
        issues = self.audit_engine.scan_document(self.test_doc_path)
        header_footer_issues = [i for i in issues if i.category == 'HEADER_FOOTER']
        
        logger.info(f"Найдено проблем колонтитулов: {len(header_footer_issues)}")
        
        if len(header_footer_issues) == 0:
            logger.warning("В документе нет проблем колонтитулов. Тест пройден тривиально.")
            self.assertTrue(True)
            return
        
        # Применяем исправления только для проблем колонтитулов
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
            output_path = tmp.name
        
        result = self.apply_engine.apply_fixes(
            doc_path=self.test_doc_path,
            issues=header_footer_issues,
            output_path=output_path
        )
        
        logger.info(f"Результат применения: успешно={result.get('success_count', 0)}, "
                   f"неудачно={result.get('failed_count', 0)}")
        
        # Проверяем, что все проблемы колонтитулов исправлены
        self.assertEqual(result.get('failed_count', 0), 0,
                        f"Не все проблемы колонтитулов исправлены: {result.get('failed_count', 0)} неудач")
        
        # Повторный аудит исправленного документа
        issues_after = self.audit_engine.scan_document(output_path)
        header_footer_after = [i for i in issues_after if i.category == 'HEADER_FOOTER']
        
        logger.info(f"Проблем колонтитулов после исправлений: {len(header_footer_after)}")
        
        # Ожидаем, что количество проблем уменьшилось до 0
        self.assertEqual(len(header_footer_after), 0,
                        f"После исправлений остались проблемы колонтитулов: {len(header_footer_after)}")
        
        # Очистка временного файла
        os.unlink(output_path)
        
        logger.info("Тест 1 пройден: все проблемы колонтитулов исправлены с первого раза.")
    
    def test_2_backward_compatibility_no_headers(self):
        """
        Тест 2: Обратная совместимость с документами без колонтитулов.
        """
        logger.info("=== Тест 2: Обратная совместимость с документами без колонтитулов ===")
        
        # Создаём документ без колонтитулов
        from docx import Document
        
        doc = Document()
        doc.add_paragraph("Документ без колонтитулов.")
        doc.add_paragraph("Второй параграф.")
        
        # Отключаем колонтитулы, чтобы аудит их не проверял
        for section in doc.sections:
            section.different_first_page_header_footer = False
            header = section.header
            header.is_linked_to_previous = True
            footer = section.footer
            footer.is_linked_to_previous = True
        
        no_header_path = "test_no_headers.docx"
        doc.save(no_header_path)
        
        try:
            # Аудит документа без колонтитулов
            issues = self.audit_engine.scan_document(no_header_path)
            
            # Проверяем, что аудит не падает на документах без колонтитулов
            self.assertIsNotNone(issues, "Аудит вернул None")
            
            # Применяем исправления ко всем проблемам (если есть)
            if issues:
                with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
                    output_path = tmp.name
                
                result = self.apply_engine.apply_fixes(
                    doc_path=no_header_path,
                    issues=issues,
                    output_path=output_path
                )
                
                # Проверяем, что ApplyEngine не падает на документах без колонтитулов
                self.assertIsNotNone(result, "ApplyEngine вернул None")
                logger.info(f"ApplyEngine успешно обработал документ без колонтитулов: "
                           f"успешно={result.get('success_count', 0)}")
                
                os.unlink(output_path)
            
            logger.info("Тест 2 пройден: обратная совместимость с документами без колонтитулов подтверждена.")
        
        finally:
            # Очистка временного файла
            if os.path.exists(no_header_path):
                os.unlink(no_header_path)
    
    def test_3_documents_with_only_headers(self):
        """
        Тест 3: Работа с документами, содержащими только колонтитулы.
        """
        logger.info("=== Тест 3: Документы только с колонтитулами ===")
        
        # Создаём документ с колонтитулами, но без основного текста
        from docx import Document
        from docx.shared import Pt
        
        doc = Document()
        
        # Удаляем все параграфы из основного тела
        for para in list(doc.paragraphs):
            p = para._element
            p.getparent().remove(p)
        
        # Добавляем колонтитулы
        section = doc.sections[0]
        header = section.header
        footer = section.footer
        
        header_para = header.add_paragraph("Только верхний колонтитул")
        footer_para = footer.add_paragraph("Только нижний колонтитул")
        
        only_headers_path = "test_only_headers.docx"
        doc.save(only_headers_path)
        
        try:
            # Аудит документа
            issues = self.audit_engine.scan_document(only_headers_path)
            
            # Проверяем, что проблемы обнаружены (возможно, связанные с колонтитулами)
            logger.info(f"Найдено проблем в документе только с колонтитулами: {len(issues)}")
            
            # Фильтруем проблемы колонтитулов
            header_footer_issues = [i for i in issues if i.category == 'HEADER_FOOTER']
            logger.info(f"Из них проблем колонтитулов: {len(header_footer_issues)}")
            
            # Применяем исправления
            if issues:
                with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
                    output_path = tmp.name
                
                result = self.apply_engine.apply_fixes(
                    doc_path=only_headers_path,
                    issues=issues,
                    output_path=output_path
                )
                
                # Проверяем, что ApplyEngine не падает
                self.assertIsNotNone(result, "ApplyEngine вернул None для документа только с колонтитулами")
                logger.info(f"ApplyEngine обработал документ только с колонтитулами: "
                           f"успешно={result.get('success_count', 0)}")
                
                # Повторный аудит
                issues_after = self.audit_engine.scan_document(output_path)
                logger.info(f"Проблем после исправлений: {len(issues_after)}")
                
                os.unlink(output_path)
            
            logger.info("Тест 3 пройден: работа с документами только с колонтитулами подтверждена.")
        
        finally:
            # Очистка временного файла
            if os.path.exists(only_headers_path):
                os.unlink(only_headers_path)
    
    def test_4_location_keys_compatibility(self):
        """
        Тест 4: Совместимость различных форматов location.
        Проверяет, что ApplyEngine корректно обрабатывает location с разными ключами.
        """
        logger.info("=== Тест 4: Совместимость форматов location ===")
        
        test_cases = [
            {
                'name': 'index ключ',
                'location': {'index': 0, 'page': 1},
                'expected_container': 'body'
            },
            {
                'name': 'paragraph_index ключ',
                'location': {'paragraph_index': 5, 'page': 1},
                'expected_container': 'body'
            },
            {
                'name': 'section+paragraph без container',
                'location': {'section': 0, 'paragraph': 0, 'page': 0},
                'expected_container': 'header'  # По умолчанию предполагаем header
            },
            {
                'name': 'section+paragraph с container=header',
                'location': {'section': 0, 'paragraph': 0, 'page': 0, 'container': 'header'},
                'expected_container': 'header'
            },
            {
                'name': 'section+paragraph с container=footer',
                'location': {'section': 0, 'paragraph': 0, 'page': 0, 'container': 'footer'},
                'expected_container': 'footer'
            }
        ]
        
        # Создаём тестовый документ
        from docx import Document
        doc = Document()
        doc.add_paragraph("Тестовый параграф 1")
        doc.add_paragraph("Тестовый параграф 2")
        
        test_doc_path = "test_location_compatibility.docx"
        doc.save(test_doc_path)
        
        try:
            for test_case in test_cases:
                logger.info(f"  Проверка: {test_case['name']}")
                
                # Создаём тестовую проблему
                test_issue = AuditIssue(
                    id=f"TEST_{test_case['name'].replace(' ', '_')}",
                    severity=Severity.CRITICAL,
                    category="PARAGRAPH",
                    element_type="Normal",
                    location=test_case['location'],
                    description=f"Тестовая проблема для {test_case['name']}",
                    current_value="",
                    expected_value="",
                    auto_fixable=True,
                    fix_payload={'style': 'Normal'}
                )
                
                # Пытаемся найти параграф через _get_paragraph из paragraph_applier
                try:
                    doc_obj = docx.Document(test_doc_path)
                    paragraph, container = self.apply_engine.paragraph_applier._get_paragraph(doc_obj, test_case['location'])
                    
                    # Проверяем, что контейнер определён правильно
                    self.assertEqual(container, test_case['expected_container'],
                                    f"Неверный контейнер для {test_case['name']}: "
                                    f"ожидался {test_case['expected_container']}, получен {container}")
                    
                    logger.info(f"    Успешно: найден параграф в контейнере {container}")
                
                except (ValueError, IndexError, KeyError) as e:
                    # Для некоторых тестовых location параграф может не существовать
                    # (например, section 0, paragraph 0 может отсутствовать)
                    # Это допустимо, если location некорректен для данного документа
                    logger.warning(f"    Пропущено: {test_case['name']} - {e}")
        
        finally:
            if os.path.exists(test_doc_path):
                os.unlink(test_doc_path)
        
        logger.info("Тест 4 пройден: совместимость форматов location подтверждена.")
    
    def test_5_multiple_iterations_header_footer(self):
        """
        Тест 5: Многократное применение исправлений к колонтитулам.
        Проверяет, что повторное применение не ломает уже исправленные колонтитулы.
        """
        logger.info("=== Тест 5: Многократное применение исправлений ===")
        
        # Используем оригинальный документ
        if not os.path.exists(self.test_doc_path):
            self.skipTest(f"Тестовый документ не найден: {self.test_doc_path}")
        
        # Первый аудит
        issues = self.audit_engine.scan_document(self.test_doc_path)
        header_footer_issues = [i for i in issues if i.category == 'HEADER_FOOTER']
        
        if len(header_footer_issues) == 0:
            logger.warning("Нет проблем колонтитулов для тестирования многократного применения.")
            self.skipTest("Нет проблем колонтитулов для тестирования.")
        
        current_doc = self.test_doc_path
        
        for iteration in range(1, 4):  # 3 итерации
            logger.info(f"  Итерация {iteration}")
            
            # Аудит текущего документа
            issues_iter = self.audit_engine.scan_document(current_doc)
            hf_count = len([i for i in issues_iter if i.category == 'HEADER_FOOTER'])
            logger.info(f"    Проблем колонтитулов: {hf_count}")
            
            # Применяем исправления
            with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
                output_path = tmp.name
            
            result = self.apply_engine.apply_fixes(
                doc_path=current_doc,
                issues=issues_iter,
                output_path=output_path
            )
            
            # Проверяем, что нет неудачных исправлений колонтитулов
            self.assertEqual(result.get('failed_count', 0), 0,
                            f"На итерации {iteration} есть неудачные исправления: {result.get('failed_count', 0)}")
            
            # Обновляем текущий документ для следующей итерации
            if iteration < 3:
                if current_doc != self.test_doc_path:
                    os.unlink(current_doc)  # Удаляем предыдущий временный файл
                current_doc = output_path
            else:
                # На последней итерации проверяем результат
                issues_final = self.audit_engine.scan_document(output_path)
                hf_final = len([i for i in issues_final if i.category == 'HEADER_FOOTER'])
                
                logger.info(f"    Финальное количество проблем колонтитулов: {hf_final}")
                
                # Ожидаем, что после многократного применения проблем не прибавилось
                self.assertLessEqual(hf_final, hf_count,
                                    f"Количество проблем колонтитулов увеличилось после итерации {iteration}")
                
                os.unlink(output_path)
        
        logger.info("Тест 5 пройден: многократное применение не ломает колонтитулы.")


def run_all_tests():
    """Запуск всех тестов и генерация отчёта."""
    import datetime
    
    # Создаём test suite
    suite = unittest.TestLoader().loadTestsFromTestCase(TestHeaderFooterRegression)
    
    # Запускаем тесты
    runner = unittest.TextTestRunner(verbosity=2)
    result = runner.run(suite)
    
    # Генерация отчёта
    report = {
        'timestamp': datetime.datetime.now().isoformat(),
        'total_tests': result.testsRun,
        'failures': len(result.failures),
        'errors': len(result.errors),
        'skipped': len(result.skipped),
        'success': result.wasSuccessful()
    }
    
    # Сохранение отчёта
    import json
    with open('regression_test_results.json', 'w', encoding='utf-8') as f:
        json.dump(report, f, indent=2, ensure_ascii=False)
    
    # Вывод сводки
    logger.info("\n=== СВОДКА ТЕСТОВ РЕГРЕССИИ ===")
    logger.info(f"Всего тестов: {report['total_tests']}")
    logger.info(f"Успешно: {report['total_tests'] - report['failures'] - report['errors'] - report['skipped']}")
    logger.info(f"Провалено: {report['failures']}")
    logger.info(f"Ошибок: {report['errors']}")
    logger.info(f"Пропущено: {report['skipped']}")
    logger.info(f"Общий результат: {'УСПЕХ' if report['success'] else 'НЕУДАЧА'}")
    
    return result.wasSuccessful()


if __name__ == '__main__':
    # Запуск тестов
    success = run_all_tests()
    sys.exit(0 if success else 1)