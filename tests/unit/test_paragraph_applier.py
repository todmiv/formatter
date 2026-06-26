#!/usr/bin/env python3
"""Тесты для ParagraphApplier и _is_spacer_paragraph."""

import os
import tempfile
import shutil
import pytest
from unittest.mock import Mock, patch, MagicMock
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from src.apply.paragraph_applier import ParagraphApplier, _is_spacer_paragraph
from src.core.audit_engine import AuditIssue, Severity
from src.core.config_loader import ConfigLoader


class TestIsSpacerParagraph:
    """Тесты функции _is_spacer_paragraph."""

    def test_empty_paragraph_is_not_spacer(self):
        doc = Document()
        para = doc.add_paragraph('')
        assert _is_spacer_paragraph(para) is False

    def test_paragraph_with_text_is_not_spacer(self):
        doc = Document()
        para = doc.add_paragraph('Text')
        assert _is_spacer_paragraph(para) is False

    def test_paragraph_with_exact_small_line_is_spacer(self):
        doc = Document()
        para = doc.add_paragraph('')
        pPr = para._element.get_or_add_pPr()
        spacing = OxmlElement('w:spacing')
        spacing.set(qn('w:lineRule'), 'exact')
        spacing.set(qn('w:line'), '20')
        pPr.append(spacing)
        assert _is_spacer_paragraph(para) is True

    def test_paragraph_with_exact_large_line_is_not_spacer(self):
        doc = Document()
        para = doc.add_paragraph('')
        pPr = para._element.get_or_add_pPr()
        spacing = OxmlElement('w:spacing')
        spacing.set(qn('w:lineRule'), 'exact')
        spacing.set(qn('w:line'), '100')
        pPr.append(spacing)
        assert _is_spacer_paragraph(para) is False

    def test_paragraph_with_auto_line_is_not_spacer(self):
        doc = Document()
        para = doc.add_paragraph('')
        pPr = para._element.get_or_add_pPr()
        spacing = OxmlElement('w:spacing')
        spacing.set(qn('w:lineRule'), 'auto')
        spacing.set(qn('w:line'), '10')
        pPr.append(spacing)
        assert _is_spacer_paragraph(para) is False


class TestParagraphApplier:
    """Тесты для ParagraphApplier."""

    def setup_method(self):
        self.temp_dir = tempfile.mkdtemp()
        self.config_path = os.path.join(self.temp_dir, 'config.yaml')
        with open(self.config_path, 'w') as f:
            f.write("""
styles:
  Normal:
    enabled: true
    font:
      name: "Times New Roman"
      size: 12
      bold: false
      italic: false
    paragraph:
      first_line_indent_cm: 1.25
      line_spacing: 1.15
      space_before_pt: 0
      space_after_pt: 0
      alignment: "justify"
detection_rules: {}
""")
        self.config = ConfigLoader(self.config_path)
        self.config.load()
        self.applier = ParagraphApplier(self.config)

    def teardown_method(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_apply_success(self):
        doc = Document()
        doc.add_paragraph('Test paragraph')

        issue = AuditIssue(
            id='TEST_1',
            severity=Severity.WARNING,
            category='PARAGRAPH',
            element_type='Normal',
            location={'index': 0},
            description='Test issue',
            current_value='',
            expected_value='',
            auto_fixable=True,
        )

        style_cache = {}
        result = self.applier.apply(doc, issue, style_cache)

        assert result is True

    def test_apply_missing_style_config(self):
        doc = Document()
        doc.add_paragraph('Test paragraph')

        issue = AuditIssue(
            id='TEST_1',
            severity=Severity.WARNING,
            category='PARAGRAPH',
            element_type='NonExistentStyle',
            location={'index': 0},
            description='Test issue',
            current_value='',
            expected_value='',
            auto_fixable=True,
        )

        style_cache = {}
        result = self.applier.apply(doc, issue, style_cache)

        assert result is False

    def test_apply_with_clear_formatting(self):
        doc = Document()
        para = doc.add_paragraph('Test')
        for run in para.runs:
            run.font.bold = True

        issue = AuditIssue(
            id='TEST_1',
            severity=Severity.WARNING,
            category='PARAGRAPH',
            element_type='Normal',
            location={'index': 0},
            description='Test issue',
            current_value='',
            expected_value='',
            auto_fixable=True,
        )

        style_cache = {}
        result = self.applier.apply(
            doc, issue, style_cache,
            clear_direct_formatting=True
        )

        assert result is True

    def test_apply_without_styles(self):
        doc = Document()
        doc.add_paragraph('Test paragraph')

        issue = AuditIssue(
            id='TEST_1',
            severity=Severity.WARNING,
            category='PARAGRAPH',
            element_type='Normal',
            location={'index': 0},
            description='Test issue',
            current_value='',
            expected_value='',
            auto_fixable=True,
        )

        style_cache = {}
        result = self.applier.apply(
            doc, issue, style_cache,
            apply_styles=False,
            apply_direct_overrides=True
        )

        assert result is True

    def test_apply_with_context(self):
        doc = Document()
        doc.add_paragraph('Para 1')
        doc.add_paragraph('Para 2')
        doc.add_paragraph('Para 3')

        issues = [
            AuditIssue(
                id=f'TEST_{i}',
                severity=Severity.WARNING,
                category='PARAGRAPH',
                element_type='Normal',
                location={'index': i},
                description=f'Test issue {i}',
                current_value='',
                expected_value='',
                auto_fixable=True,
            )
            for i in range(3)
        ]

        style_cache = {}
        result = self.applier.apply_with_context(doc, issues, style_cache)

        assert result['applied'] == 3
        assert result['failed'] == 0

    def test_apply_with_context_non_paragraph_issues(self):
        doc = Document()
        doc.add_paragraph('Para 1')

        issues = [
            AuditIssue(
                id='TABLE_1',
                severity=Severity.WARNING,
                category='TABLE',
                element_type='Table Grid',
                location={'table_index': 0},
                description='Table issue',
                current_value='',
                expected_value='',
                auto_fixable=True,
            )
        ]

        style_cache = {}
        result = self.applier.apply_with_context(doc, issues, style_cache)

        assert result['applied'] == 0
        assert result['failed'] == 0
