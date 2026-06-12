#!/usr/bin/env python3
"""
Финальный тестовый цикл для достижения 100% эффективности исправлений.
Загружает исходный документ, запускает аудит, применяет исправления с оптимизацией,
повторяет цикл до достижения 0 проблем или максимум 5 итераций.
Создаёт отчёт с детальной статистикой.
"""

import sys
import os
import json
import logging
import time
from typing import List, Dict, Any, Optional
from datetime import datetime

# Добавляем текущую директорию в путь для импорта модулей
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from src.core.config_loader import ConfigLoader
from src.core.audit_engine import AuditEngine, AuditIssue
from src.apply.apply_orchestrator import ApplyOrchestrator
from src.core.optimized_application import OptimizedApplyEngine

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

class HundredPercentEfficiencyTest:
    """Тестовый цикл для достижения 100% эффективности."""
    
    def __init__(self, config_path: str = "config_v4.2.yaml"):
        self.config_path = config_path
        self.config = ConfigLoader(config_path)
        self.config.load()
        self.audit_engine = AuditEngine(self.config)
        self.apply_engine = ApplyOrchestrator(self.config)
        self.optimized_engine = OptimizedApplyEngine(self.config)
        
        self.results = {
            'start_time': None,
            'end_time': None,
            'document': None,
            'initial_issues': 0,
            'iterations': [],
            'final_issues': 0,
            'efficiency': 0.0,
            'total_time_seconds': 0,
            'success': False
        }
    
    def run_test(self, doc_path: str, max_iterations: int = 5) -> Dict[str, Any]:
        """
        Запускает тестовый цикл.
        
        :param doc_path: Путь к исходному документу.
        :param max_iterations: Максимальное количество итераций.
        :return: Результаты теста.
        """
        if not os.path.exists(doc_path):
            raise FileNotFoundError(f"Документ не найден: {doc_path}")
        
        self.results['start_time'] = datetime.now().isoformat()
        self.results['document'] = doc_path
        
        logger.info(f"Начало теста 100% эффективности для документа: {doc_path}")
        logger.info(f"Максимальное количество итераций: {max_iterations}")
        
        current_doc = doc_path
        iteration = 0
        
        while iteration < max_iterations:
            iteration += 1
            logger.info(f"\n{'='*60}")
            logger.info(f"Итерация {iteration}")
            logger.info(f"{'='*60}")
            
            # 1. Запускаем аудит
            logger.info("Запуск аудита...")
            issues = self.audit_engine.scan_document(current_doc)
            issues_count = len(issues)
            
            if iteration == 1:
                self.results['initial_issues'] = issues_count
            
            logger.info(f"Найдено проблем: {issues_count}")
            
            # Если проблем нет, завершаем
            if issues_count == 0:
                logger.info("Достигнуто 100% эффективности! Проблем не осталось.")
                self.results['success'] = True
                break
            
            # 2. Анализируем зависимости
            logger.info("Анализ зависимостей...")
            dependency_analysis = self.optimized_engine.analyze_dependencies(issues)
            logger.info(f"  Групп: {dependency_analysis['group_count']}")
            logger.info(f"  Циклов: {dependency_analysis['cycles_found']}")
            
            # 3. Применяем исправления с оптимизацией
            logger.info("Применение исправлений с оптимизацией...")
            output_path = current_doc.replace('.docx', f'_iter{iteration}.docx')
            
            apply_results = self.optimized_engine.apply_optimized(
                current_doc, issues, output_path,
                max_iterations=3,  # Внутренние итерации
                apply_styles=True,
                apply_direct_overrides=True,
                clear_direct_formatting=False
            )
            
            # 4. Запускаем повторный аудит для проверки результата
            logger.info("Проверка результата...")
            remaining_issues = self.audit_engine.scan_document(output_path)
            remaining_count = len(remaining_issues)
            
            # Собираем статистику по итерации
            iteration_stats = {
                'iteration': iteration,
                'issues_before': issues_count,
                'issues_after': remaining_count,
                'applied': apply_results['total_applied'],
                'failed': apply_results['total_failed'],
                'groups_processed': len(apply_results.get('group_results', [])),
                'dependency_groups': dependency_analysis['group_count'],
                'cycles': dependency_analysis['cycles_found'],
                'output_document': output_path,
                'remaining_by_category': self._count_by_category(remaining_issues)
            }
            
            self.results['iterations'].append(iteration_stats)
            
            logger.info(f"Результат итерации {iteration}:")
            logger.info(f"  Проблем до: {issues_count}")
            logger.info(f"  Проблем после: {remaining_count}")
            logger.info(f"  Устранено: {issues_count - remaining_count}")
            logger.info(f"  Эффективность итерации: {(issues_count - remaining_count) / issues_count * 100:.1f}%")
            
            # Если прогресс отсутствует, останавливаемся
            if remaining_count >= issues_count:
                logger.warning("Прогресс отсутствует. Прерываем цикл.")
                break
            
            # Обновляем текущий документ для следующей итерации
            current_doc = output_path
        
        # Финальный аудит
        logger.info("\n" + "="*60)
        logger.info("Финальный аудит...")
        final_issues = self.audit_engine.scan_document(current_doc)
        self.results['final_issues'] = len(final_issues)
        self.results['final_document'] = current_doc
        
        # Рассчитываем общую эффективность
        if self.results['initial_issues'] > 0:
            self.results['efficiency'] = 1.0 - (self.results['final_issues'] / self.results['initial_issues'])
        else:
            self.results['efficiency'] = 1.0
        
        self.results['success'] = (self.results['final_issues'] == 0)
        self.results['end_time'] = datetime.now().isoformat()
        
        # Рассчитываем общее время
        start = datetime.fromisoformat(self.results['start_time'])
        end = datetime.fromisoformat(self.results['end_time'])
        self.results['total_time_seconds'] = (end - start).total_seconds()
        
        # Выводим итоговый отчёт
        self._print_summary()
        
        # Сохраняем результаты
        self._save_results()
        
        return self.results
    
    def _count_by_category(self, issues: List[AuditIssue]) -> Dict[str, int]:
        """Считает проблемы по категориям."""
        counts = {}
        for issue in issues:
            cat = issue.category
            counts[cat] = counts.get(cat, 0) + 1
        return counts
    
    def _print_summary(self):
        """Выводит итоговый отчёт в консоль."""
        print("\n" + "="*80)
        print("ИТОГОВЫЙ ОТЧЁТ ТЕСТА 100% ЭФФЕКТИВНОСТИ")
        print("="*80)
        print(f"Документ: {self.results['document']}")
        print(f"Начальное количество проблем: {self.results['initial_issues']}")
        print(f"Финальное количество проблем: {self.results['final_issues']}")
        print(f"Общая эффективность: {self.results['efficiency']:.1%}")
        print(f"Успех (0 проблем): {'ДА' if self.results['success'] else 'НЕТ'}")
        print(f"Всего итераций: {len(self.results['iterations'])}")
        print(f"Общее время: {self.results['total_time_seconds']:.1f} секунд")
        print(f"Финальный документ: {self.results.get('final_document', 'N/A')}")
        
        print("\nДетали по итерациям:")
        for iter_data in self.results['iterations']:
            print(f"  Итерация {iter_data['iteration']}:")
            print(f"    Проблем до/после: {iter_data['issues_before']} → {iter_data['issues_after']}")
            print(f"    Устранено: {iter_data['issues_before'] - iter_data['issues_after']}")
            print(f"    Применено исправлений: {iter_data['applied']}")
            print(f"    Групп зависимостей: {iter_data['dependency_groups']}")
        
        # Анализ оставшихся проблем
        if self.results['final_issues'] > 0:
            print("\nОСТАВШИЕСЯ ПРОБЛЕМЫ:")
            final_issues = self.audit_engine.scan_document(self.results.get('final_document', self.results['document']))
            by_category = self._count_by_category(final_issues)
            for cat, count in by_category.items():
                print(f"  {cat}: {count}")
        
        print("="*80)
    
    def _save_results(self, output_path: str = "100_percent_efficiency_results.json"):
        """Сохраняет результаты теста в JSON файл."""
        # Конвертируем объекты AuditIssue в словари для сериализации
        serializable_results = self.results.copy()
        
        # Добавляем детальную информацию о финальных проблемах
        if self.results.get('final_document'):
            final_issues = self.audit_engine.scan_document(self.results['final_document'])
            serializable_results['final_issues_details'] = [
                {
                    'id': issue.id,
                    'category': issue.category,
                    'element_type': issue.element_type,
                    'description': issue.description,
                    'severity': issue.severity.value,
                    'location': issue.location,
                    'expected_value': issue.expected_value,
                    'current_value': issue.current_value
                }
                for issue in final_issues
            ]
        
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(serializable_results, f, ensure_ascii=False, indent=2)
        
        logger.info(f"Результаты сохранены в {output_path}")
    
    def validate_results(self) -> Dict[str, Any]:
        """
        Проводит валидацию результатов.
        Убеждается, что все проблемы действительно устранены, а не скрыты.
        """
        validation = {
            'document_structure_preserved': True,
            'styles_applied_correctly': True,
            'backward_compatibility': True,
            'issues_truly_fixed': True,
            'warnings': []
        }
        
        if not self.results.get('final_document'):
            validation['warnings'].append('Финальный документ не создан')
            return validation
        
        final_doc = self.results['final_document']
        
        # 1. Проверяем, что документ открывается
        try:
            import docx
            doc = docx.Document(final_doc)
            validation['document_structure_preserved'] = True
        except Exception as e:
            validation['document_structure_preserved'] = False
            validation['warnings'].append(f'Ошибка открытия документа: {e}')
        
        # 2. Проверяем, что стили применены правильно
        # (упрощённая проверка - можно расширить)
        try:
            doc = docx.Document(final_doc)
            # Проверяем наличие основных стилей
            required_styles = ['Normal', 'Heading 1', 'Heading 2', 'Heading 3']
            for style_name in required_styles:
                try:
                    style = doc.styles[style_name]
                except KeyError:
                    validation['styles_applied_correctly'] = False
                    validation['warnings'].append(f'Стиль {style_name} отсутствует')
        except Exception as e:
            validation['warnings'].append(f'Ошибка проверки стилей: {e}')
        
        # 3. Проверяем обратную совместимость
        # (упрощённо - проверяем, что документ можно сохранить и открыть)
        try:
            temp_path = final_doc.replace('.docx', '_compatibility_test.docx')
            doc.save(temp_path)
            test_doc = docx.Document(temp_path)
            os.remove(temp_path)
            validation['backward_compatibility'] = True
        except Exception as e:
            validation['backward_compatibility'] = False
            validation['warnings'].append(f'Ошибка совместимости: {e}')
        
        # 4. Проверяем, что проблемы действительно устранены
        # Запускаем аудит с нулевыми допусками для строгой проверки
        strict_audit = AuditEngine(self.config, font_size_tolerance_twips=0,
                                  indent_tolerance_twips=0, spacing_tolerance_twips=0)
        strict_issues = strict_audit.scan_document(final_doc)
        
        if len(strict_issues) > 0:
            validation['issues_truly_fixed'] = False
            validation['warnings'].append(f'Найдено {len(strict_issues)} проблем при строгой проверке')
            validation['strict_issues_count'] = len(strict_issues)
        
        return validation

def main():
    """Основная функция скрипта."""
    import argparse
    
    parser = argparse.ArgumentParser(description='Тест 100% эффективности исправлений')
    parser.add_argument('--document', '-d', default='Том II_Волжск.docx',
                       help='Путь к исходному документу (по умолчанию: Том II_Волжск.docx)')
    parser.add_argument('--iterations', '-i', type=int, default=5,
                       help='Максимальное количество итераций (по умолчанию: 5)')
    parser.add_argument('--config', '-c', default='config_v4.2.yaml',
                       help='Путь к конфигурационному файлу')
    parser.add_argument('--validate', '-v', action='store_true',
                       help='Провести валидацию результатов')
    
    args = parser.parse_args()
    
    # Проверяем существование документа
    if not os.path.exists(args.document):
        logger.error(f"Документ не найден: {args.document}")
        logger.info("Доступные документы:")
        for f in os.listdir('.'):
            if f.endswith('.docx'):
                logger.info(f"  - {f}")
        sys.exit(1)
    
    # Запускаем тест
    test = HundredPercentEfficiencyTest(args.config)
    results = test.run_test(args.document, max_iterations=args.iterations)
    
    # Проводим валидацию, если запрошено
    if args.validate:
        logger.info("\nПроведение валидации результатов...")
        validation = test.validate_results()
        
        print("\n" + "="*80)
        print("РЕЗУЛЬТАТЫ ВАЛИДАЦИИ")
        print("="*80)
        for key, value in validation.items():
            if key != 'warnings':
                print(f"{key}: {'ПРОЙДЕНО' if value else 'НЕ ПРОЙДЕНО'}")
        
        if validation['warnings']:
            print("\nПредупреждения:")
            for warning in validation['warnings']:
                print(f"  - {warning}")
        
        # Сохраняем результаты валидации
        with open('validation_results.json', 'w', encoding='utf-8') as f:
            json.dump(validation, f, ensure_ascii=False, indent=2)
        
        print("="*80)
    
    # Вывод итогового статуса
    if results['success']:
        logger.info("\n🎉 ТЕСТ ПРОЙДЕН УСПЕШНО! Достигнуто 100% эффективности исправлений.")
        sys.exit(0)
    else:
        logger.warning("\n⚠️ ТЕСТ НЕ ПРОЙДЕН. Остались неисправленные проблемы.")
        sys.exit(1)

if __name__ == "__main__":
    main()