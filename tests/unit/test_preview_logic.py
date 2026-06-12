import sys
import os
sys.path.insert(0, os.path.dirname(__file__))

import docx
from src.core.audit_engine import AuditEngine, Severity
from src.core.config_loader import ConfigLoader

def test_preview_logic():
    # Используем существующий тестовый файл
    doc_path = 'test_output.docx'
    if not os.path.exists(doc_path):
        # Попробуем найти другой
        docx_files = [f for f in os.listdir('.') if f.lower().endswith('.docx')]
        if docx_files:
            doc_path = docx_files[0]
        else:
            print("Нет DOCX файлов для теста.")
            return
    config_path = 'configs/active/config.yaml'
    
    loader = ConfigLoader(config_path)
    loader.load()
    engine = AuditEngine(loader)
    issues = engine.scan_document(doc_path)
    
    if not issues:
        print("Нет проблем для теста предпросмотра.")
        return
    
    for issue in issues[:3]:  # первые три проблемы
        para_index = issue.location.get('index')
        print(f"\nПроблема: {issue.id}, индекс: {para_index}")
        
        preview = ""
        if para_index is not None and doc_path and os.path.exists(doc_path):
            try:
                doc = docx.Document(doc_path)
                if 0 <= para_index < len(doc.paragraphs):
                    para = doc.paragraphs[para_index]
                    preview = para.text.strip()
                    if not preview:
                        preview = "(пустой параграф)"
                    # Добавим контекст: предыдущий и следующий параграфы
                    context = []
                    if para_index > 0:
                        prev = doc.paragraphs[para_index - 1].text.strip()
                        if prev:
                            context.append(f"← {prev[:100]}...")
                    if para_index < len(doc.paragraphs) - 1:
                        nxt = doc.paragraphs[para_index + 1].text.strip()
                        if nxt:
                            context.append(f"→ {nxt[:100]}...")
                    if context:
                        preview += "\n\nКонтекст:\n" + "\n".join(context)
                else:
                    preview = f"Параграф #{para_index} не найден в документе."
            except Exception as e:
                preview = f"Ошибка при извлечении фрагмента: {e}"
        else:
            preview = f"Фрагмент текста с ошибкой не доступен.\nПараграф #{para_index if para_index is not None else '?'}"
        
        print("Предпросмотр:")
        print(preview)
        if "Фрагмент текста с ошибкой не доступен" in preview:
            print("ВНИМАНИЕ: всё ещё показывает заглушку!")
        else:
            print("Успех: заглушка заменена реальным текстом.")

if __name__ == "__main__":
    test_preview_logic()