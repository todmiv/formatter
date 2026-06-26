#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
DocumentTreeMixin — работа с деревом документа.
Вынесена из GOSTFormatterGUI для уменьшения размера класса.
"""

import threading
import tkinter as tk
from tkinter import messagebox
from typing import Dict, Any


class DocumentTreeMixin:
    """
    Миксин для работы с деревом документа.
    Требует атрибутов: controller, doc_tree, log, root.
    """

    def populate_doc_tree(self, doc_path: str):
        """Заполнение дерева документа в отдельном потоке."""
        thread = threading.Thread(
            target=self._populate_doc_tree_thread,
            args=(doc_path,),
            daemon=True
        )
        thread.start()

    def _populate_doc_tree_thread(self, doc_path: str):
        """Поток заполнения дерева."""
        try:
            data = self._collect_doc_tree_data(doc_path)
            self.root.after(0, lambda: self._insert_doc_tree_data(data))
        except Exception as e:
            self.root.after(0, lambda: self._on_populate_error(e))

    def _collect_doc_tree_data(self, doc_path: str) -> Dict[str, Any]:
        """Сбор данных для дерева документа."""
        import docx

        if not doc_path or not hasattr(self.controller, 'doc_path'):
            return {'items': [], 'counts': {}}

        doc = docx.Document(doc_path)

        data = {
            'filename': doc_path.split('/')[-1].split('\\')[-1],
            'items': [],
            'counts': {
                'heading': 0,
                'table': 0,
                'list': 0,
                'para': 0,
            }
        }

        for idx, para in enumerate(doc.paragraphs):
            style_name = para.style.name
            text = para.text[:50] if para.text else ''

            if style_name.startswith('Heading'):
                level = int(style_name.split()[-1]) if style_name.split()[-1].isdigit() else 1
                data['items'].append({
                    'type': 'heading',
                    'level': level,
                    'text': text,
                    'index': idx,
                })
                data['counts']['heading'] += 1
            elif style_name == 'List Bullet' or style_name == 'List Number':
                data['items'].append({
                    'type': 'list',
                    'text': text,
                    'index': idx,
                })
                data['counts']['list'] += 1
            else:
                data['items'].append({
                    'type': 'paragraph',
                    'text': text,
                    'index': idx,
                })
                data['counts']['para'] += 1

        for idx, table in enumerate(doc.tables):
            rows = len(table.rows)
            cols = len(table.columns)
            data['items'].append({
                'type': 'table',
                'text': f'Таблица {idx + 1} ({rows}x{cols})',
                'index': idx,
            })
            data['counts']['table'] += 1

        return data

    def _insert_doc_tree_data(self, data: Dict[str, Any]):
        """Вставка данных в дерево документа."""
        if not hasattr(self, 'doc_tree'):
            return

        self.doc_tree.delete(*self.doc_tree.get_children())

        root_id = self.doc_tree.insert('', 'end', text=data.get('filename', 'Документ'), open=True)

        counts = data.get('counts', {})
        summary = f"З: {counts.get('heading', 0)} | Т: {counts.get('table', 0)} | П: {counts.get('para', 0)}"
        self.doc_tree.insert(root_id, 'end', text=summary, values=('', ''))

        for item in data.get('items', []):
            item_type = item.get('type', '')
            text = item.get('text', '')
            index = item.get('index', '')

            if item_type == 'heading':
                level = item.get('level', 1)
                prefix = '  ' * (level - 1)
                self.doc_tree.insert(root_id, 'end', text=f"{prefix}H{level}: {text}", values=('heading', index))
            elif item_type == 'table':
                self.doc_tree.insert(root_id, 'end', text=f"📊 {text}", values=('table', index))
            elif item_type == 'list':
                self.doc_tree.insert(root_id, 'end', text=f"• {text}", values=('list', index))
            else:
                self.doc_tree.insert(root_id, 'end', text=text[:40], values=('para', index))

    def _on_populate_error(self, error):
        """Обработка ошибки заполнения дерева."""
        self.log(f"Ошибка заполнения дерева: {error}")

    def on_tree_select(self, event):
        """Обработка выбора элемента в дереве."""
        if not hasattr(self, 'doc_tree'):
            return

        selection = self.doc_tree.selection()
        if not selection:
            return

        item = self.doc_tree.item(selection[0])
        values = item.get('values', ())

        if len(values) >= 2:
            item_type = values[0]
            index = values[1]

            if item_type and index:
                self.log(f"Выбран элемент: {item_type} #{index}")

    def treeview_sort_column(self, col, reverse):
        """Сортировка колонки таблицы."""
        if not hasattr(self, 'tree'):
            return

        data = [(self.tree.set(k, col), k) for k in self.tree.get_children('')]
        data.sort(reverse=reverse)

        for index, (val, k) in enumerate(data):
            self.tree.move(k, '', index)

        self.tree.heading(col, command=lambda: self.treeview_sort_column(col, not reverse))
