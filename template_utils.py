"""
Шаблон-подход: копирует DOCX-шаблон и заменяет содержимое, сохраняя форматирование.
"""
import shutil, os
from docx import Document
from docx.oxml.ns import qn


TEMPLATE = r'd:\project\Demographics\Наборы данных\Миасский ГО\result\ОБРАЗЕЦ_Том II. Материалы по обоснованию_Миасский.docx'


def copy_template(output_path):
    """Копирует шаблон в новый файл."""
    shutil.copy2(TEMPLATE, output_path)
    return output_path


def clear_content(doc_path):
    """Очищает содержимое документа, сохраняя стили и настройки страницы."""
    doc = Document(doc_path)
    body = doc.element.body

    # Удаляем все абзацы и таблицы из тела
    for child in list(body):
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
        if tag in ('p', 'tbl', 'sdt'):
            body.remove(child)

    doc.save(doc_path)
    return doc


def add_paragraph(doc, text, style_name=None):
    """Добавляет абзац с нужным стилем."""
    p = doc.add_paragraph(text)
    if style_name:
        try:
            p.style = doc.styles[style_name]
        except KeyError:
            pass  # стиль не найден — оставляем Normal
    return p


def add_table(doc, rows, cols, style_name='Table Grid'):
    """Добавляет таблицу."""
    table = doc.add_table(rows=rows, cols=cols, style=style_name)
    return table


def create_document(output_path, paragraphs):
    """
    Создаёт документ из шаблона с заданным содержимым.

    Args:
        output_path: путь для сохранения
        paragraphs: список кортежей (текст, стиль_или_None)
    """
    copy_template(output_path)
    doc = Document(output_path)

    # Очищаем
    body = doc.element.body
    for child in list(body):
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
        if tag in ('p', 'tbl'):
            body.remove(child)

    # Добавляем новый контент
    for text, style in paragraphs:
        add_paragraph(doc, text, style)

    doc.save(output_path)
    return output_path


# ============================================================
# ПРИМЕР ИСПОЛЬЗОВАНИЯ
# ============================================================

if __name__ == '__main__':
    output = os.path.join(os.path.dirname(__file__), 'dist', 'templates', 'new_document.docx')

    content = [
        ('РАЗДЕЛ 1. НОВЫЙ ДОКУМЕНТ', '021216Раздел'),
        ('ГЛАВА 1.1 Введение', '021216Глава'),
        ('1.1.1 Общие сведения', '021216Подглава'),
        ('3.3.1 Основной раздел', 'Пункт'),
        ('Данный текст оформлен стилем "Текст доклада" с отступом красной строки 1,25 см и выравниванием по ширине.', 'Текст доклада'),
        ('1) Первый пункт списка', 'List Paragraph'),
        ('2) Второй пункт списка', 'List Paragraph'),
        ('Таблица 1.1 — Пример таблицы', 'Заголовок таблицы'),
    ]

    create_document(output, content)
    print(f'Создан: {output}')
